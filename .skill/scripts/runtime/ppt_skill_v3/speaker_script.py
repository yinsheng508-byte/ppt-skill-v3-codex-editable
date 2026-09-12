from __future__ import annotations

import hashlib
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .deliverable_naming import project_deliverable_relpaths
from .document_rendering import (
    add_page_number as _add_page_number,
    find_pdf_font_path as _find_pdf_font_path,
    p as _p,
    pdf_page_count as _pdf_page_count,
    pdf_text_probe as _pdf_text_probe,
    relative_or_absolute as _relative_or_absolute,
    require_locked_source_exists as _require_locked_source_exists,
    require_locked_source_matches_state as _require_locked_source_matches_state,
)
from .events import append_event
from .json_io import read_json, write_json
from .state import read_state, stage3_lesson_plan_required, write_state
from .time_utils import now_iso
from .validation import ValidationError, validate_speaker_script, validate_speaker_script_manifest


USER_DIR = "阶段3_逐字稿与教案输出/逐字稿"
STATE_DIR = "_state/阶段3"
SCRIPT_JSON_REL = f"{STATE_DIR}/speaker_script.json"
MANIFEST_REL = f"{STATE_DIR}/speaker_script_manifest.json"
FONT_FAMILY = "Arial Unicode MS"
INK = RGBColor(31, 41, 55)
MUTED = RGBColor(91, 105, 120)
ACCENT = RGBColor(15, 76, 117)
ACCENT_DARK = "0F4C75"


def build_speaker_script(run_dir: str | Path, script_json: str | Path) -> dict[str, Any]:
    root = Path(run_dir)
    state = read_state(root)
    if state["current_stage"] != "stage3" or state["status"] not in {
        "ready_for_stage3_script",
        "stage3_script_generated",
        "stage3_lesson_plan_generated",
        "stage3_lesson_plan_revision_requested",
    }:
        raise ValidationError("speaker script generation requires stage3 ready_for_stage3_script state")

    script = validate_speaker_script(read_json(script_json))
    if script["project_name"] != state["project_name"]:
        raise ValidationError("speaker_script.project_name does not match project state")
    if script["run_dir"] != state["run_dir"]:
        raise ValidationError("speaker_script.run_dir does not match project state")
    locked_source = _locked_presentation_source(script)
    _require_locked_source_exists(root, locked_source)
    _require_locked_source_matches_state(state, locked_source)

    relpaths = project_deliverable_relpaths(root, state=state, script=script)
    markdown_rel = relpaths["stage3_speaker_script"]
    docx_rel = relpaths["stage3_speaker_script_docx"]
    pdf_rel = relpaths["stage3_speaker_script_pdf"]

    (root / USER_DIR).mkdir(parents=True, exist_ok=True)
    (root / STATE_DIR / "pdf_conversion").mkdir(parents=True, exist_ok=True)
    (root / STATE_DIR / "logs").mkdir(parents=True, exist_ok=True)

    write_json(root / SCRIPT_JSON_REL, script)
    markdown = _render_markdown(script)
    (root / markdown_rel).write_text(markdown, encoding="utf-8")
    _render_docx(script, root / docx_rel)
    conversion = _build_pdf(root / docx_rel, root / pdf_rel, script)
    conversion["pdf_path"] = pdf_rel
    if conversion.get("paired_docx"):
        conversion["paired_docx"] = docx_rel
    if conversion.get("source") in {docx_rel, str(root / docx_rel)}:
        conversion["source"] = docx_rel
    script_character_count = _script_character_count(script)

    manifest = validate_speaker_script_manifest(
        {
            "schema_version": "2.0",
            "project_name": state["project_name"],
            "run_dir": state["run_dir"],
            "files": [
                _file_entry("Markdown 逐字稿", root, markdown_rel),
                _file_entry("Word 逐字稿", root, docx_rel),
                _file_entry("PDF 逐字稿", root, pdf_rel),
            ],
            "conversion": conversion,
            "summary": {
                "slides": len(script["slides"]),
                "script_characters": script_character_count,
                "estimated_minutes": _estimate_minutes(script_character_count),
                "pdf_pages": _pdf_page_count(root / pdf_rel),
                "pdf_text_probe": _pdf_text_probe(root / pdf_rel),
                "first_page_header_mode": "inline_top_header_no_cover_page",
                "docx_font_family": FONT_FAMILY,
                "docx_language": "zh-CN",
                "locked_source_mode": locked_source["source_mode"],
                "locked_source_path": locked_source["source_path"],
            },
            "created_at": now_iso(),
            "status": "generated",
        }
    )
    write_json(root / MANIFEST_REL, manifest)
    state["status"] = "stage3_script_generated"
    state["required_actor"] = "main_controller"
    stage3_outputs = state.setdefault("stage3_outputs", {})
    speaker_output = stage3_outputs.setdefault("speaker_script", {})
    speaker_output["required"] = True
    speaker_output["status"] = "generated"
    state.setdefault("user_artifacts", {})["stage3_speaker_script"] = markdown_rel
    state["user_artifacts"]["stage3_speaker_script_docx"] = docx_rel
    state["user_artifacts"]["stage3_speaker_script_pdf"] = pdf_rel
    state.setdefault("expected_user_paths", {})["stage3_speaker_script"] = markdown_rel
    state["expected_user_paths"]["stage3_speaker_script_docx"] = docx_rel
    state["expected_user_paths"]["stage3_speaker_script_pdf"] = pdf_rel
    state.setdefault("quality", {})["stage3"] = "pending_controller_review"
    if stage3_lesson_plan_required(state):
        state["next_required_action"] = "主控大模型检查阶段3逐字稿；K12 教案必选时继续生成阶段3教案设计、Word 和 PDF"
    else:
        state["next_required_action"] = "主控大模型检查阶段3逐字稿、Word 和 PDF；通过后记录阶段3输出完成决策并进入阶段4整理"
    write_state(root, state)
    append_event(root, "speaker_script_generated", "runtime", slides=len(script["slides"]))
    return manifest


def _locked_presentation_source(script: dict[str, Any]) -> dict[str, str]:
    basis = script["basis"]
    source = basis.get("locked_presentation_source")
    if isinstance(source, dict):
        result = {
            "source_mode": str(source["source_mode"]),
            "source_path": str(source["source_path"]),
        }
        if source.get("source_sha256"):
            result["source_sha256"] = str(source["source_sha256"])
        if source.get("confirmation_basis"):
            result["confirmation_basis"] = str(source["confirmation_basis"])
        return result
    raise ValidationError("speaker_script.basis.locked_presentation_source is required")


def _render_markdown(script: dict[str, Any]) -> str:
    talk = script["talk"]
    lines = [
        f"# {talk['title']}演讲逐字稿",
        "",
        f"讲者身份：{talk['speaker_role']}",
        "",
        "---",
        "",
    ]
    for slide in script["slides"]:
        lines.extend(
            [
                f"### 第 {slide['slide_index']:02d} 页｜{slide['slide_title']}",
                "",
            ]
        )
        lines.extend([f"{paragraph}" for paragraph in slide["script"]])
        lines.append("")
    lines.append("")
    return "\n".join(lines)


def _render_docx(script: dict[str, Any], target: Path) -> None:
    doc = Document()
    _configure_document(doc, script)
    _configure_styles(doc)
    _write_document_header(doc, script)
    _write_slide_scripts(doc, script)
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)


def _configure_document(doc: Document, script: dict[str, Any]) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.1)
    section.different_first_page_header_footer = True

    header = section.header.paragraphs[0]
    header.text = script["talk"]["title"]
    header.style = doc.styles["Header"]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_page_number(footer)


def _configure_styles(doc: Document) -> None:
    styles = doc.styles
    _set_style_font(styles["Normal"], FONT_FAMILY, 10.5, INK)
    styles["Normal"].paragraph_format.line_spacing = 1.4
    styles["Normal"].paragraph_format.space_after = Pt(4)

    for name, size, color in [
        ("Title", 22, RGBColor(17, 24, 39)),
        ("Subtitle", 12, MUTED),
        ("Heading 1", 15, ACCENT),
        ("Heading 2", 11, RGBColor(17, 24, 39)),
    ]:
        _set_style_font(styles[name], FONT_FAMILY, size, color, bold=name in {"Title", "Heading 1", "Heading 2"})
        styles[name].paragraph_format.keep_with_next = True
    styles["Title"].paragraph_format.space_before = Pt(0)
    styles["Title"].paragraph_format.space_after = Pt(5)
    styles["Subtitle"].paragraph_format.space_before = Pt(0)
    styles["Subtitle"].paragraph_format.space_after = Pt(4)
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(7)
    styles["Heading 2"].paragraph_format.space_before = Pt(8)
    styles["Heading 2"].paragraph_format.space_after = Pt(4)

    script_body = _paragraph_style(doc, "Script Body")
    _set_style_font(script_body, FONT_FAMILY, 11, INK)
    script_body.paragraph_format.line_spacing = 1.45
    script_body.paragraph_format.space_after = Pt(7)

    note_style = _paragraph_style(doc, "Speaker Note")
    _set_style_font(note_style, FONT_FAMILY, 9.5, MUTED)
    note_style.paragraph_format.line_spacing = 1.25
    note_style.paragraph_format.space_after = Pt(4)

    meta_style = _paragraph_style(doc, "Talk Meta")
    _set_style_font(meta_style, FONT_FAMILY, 9.5, MUTED)
    meta_style.paragraph_format.space_after = Pt(5)


def _write_document_header(doc: Document, script: dict[str, Any]) -> None:
    talk = script["talk"]
    label = doc.add_paragraph("PPT 演讲逐字稿", style="Subtitle")
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title = doc.add_paragraph(talk["title"], style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph(style="Talk Meta")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(talk["speaker_role"])
    meta.paragraph_format.space_after = Pt(6)
    _add_rule(doc, ACCENT_DARK)


def _write_slide_scripts(doc: Document, script: dict[str, Any]) -> None:
    for slide in script["slides"]:
        doc.add_heading(f"第 {slide['slide_index']:02d} 页｜{slide['slide_title']}", level=1)
        for paragraph_text in slide["script"]:
            doc.add_paragraph(paragraph_text, style="Script Body")


def _paragraph_style(doc: Document, name: str):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def _set_style_font(style, font_name: str, size_pt: float, color: RGBColor, *, bold: bool = False) -> None:
    style.font.name = font_name
    style.font.size = Pt(size_pt)
    style.font.color.rgb = color
    style.font.bold = bold
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:eastAsia"), font_name)
    r_fonts.set(qn("w:cs"), font_name)
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")


def _set_run_east_asia(run, font_name: str) -> None:
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), font_name)


def _paragraph_border(paragraph, *, left: str | None = None, bottom: str | None = None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        insert_before = {
            "w:shd",
            "w:tabs",
            "w:suppressAutoHyphens",
            "w:kinsoku",
            "w:wordWrap",
            "w:overflowPunct",
            "w:topLinePunct",
            "w:autoSpaceDE",
            "w:autoSpaceDN",
            "w:bidi",
            "w:adjustRightInd",
            "w:snapToGrid",
            "w:spacing",
            "w:ind",
            "w:contextualSpacing",
            "w:mirrorIndents",
            "w:suppressOverlap",
            "w:jc",
            "w:textDirection",
            "w:textAlignment",
            "w:textboxTightWrap",
            "w:outlineLvl",
            "w:divId",
            "w:cnfStyle",
            "w:rPr",
            "w:sectPr",
            "w:pPrChange",
        }
        for index, child in enumerate(p_pr):
            if child.tag in {qn(tag) for tag in insert_before}:
                p_pr.insert(index, p_bdr)
                break
        else:
            p_pr.append(p_bdr)
    for side, color in {"left": left, "bottom": bottom}.items():
        if not color:
            continue
        border = p_bdr.find(qn(f"w:{side}"))
        if border is None:
            border = OxmlElement(f"w:{side}")
            p_bdr.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8" if side == "left" else "6")
        border.set(qn("w:space"), "6" if side == "left" else "1")
        border.set(qn("w:color"), color)


def _add_rule(doc: Document, color: str) -> None:
    paragraph = doc.add_paragraph(style="Talk Meta")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run("-" * 64)
    run.font.name = FONT_FAMILY
    run.font.size = Pt(5)
    run.font.color.rgb = RGBColor.from_string(color)
    _set_run_east_asia(run, FONT_FAMILY)


def _build_pdf(docx_path: Path, pdf_path: Path, script: dict[str, Any]) -> dict[str, Any]:
    root = Path(script["run_dir"])
    docx_rel = _relative_or_absolute(root, docx_path)
    pdf_rel = _relative_or_absolute(root, pdf_path)
    try:
        font_path = _render_pdf_with_reportlab(script, pdf_path)
        return {
            "source": SCRIPT_JSON_REL,
            "paired_docx": docx_rel,
            "pdf_path": pdf_rel,
            "tool": "reportlab",
            "font_path": str(font_path),
        }
    except Exception as reportlab_error:
        reportlab_note = str(reportlab_error)

    external_conversion, external_error = _try_external_reportlab_pdf(pdf_path, script, reportlab_note)
    if external_conversion:
        return external_conversion
    if external_error:
        reportlab_note = f"{reportlab_note}; external ReportLab attempt failed: {external_error}"

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise ValidationError(f"ReportLab PDF generation failed and LibreOffice/soffice is not available: {reportlab_note}")
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    if pdf_path.exists():
        pdf_path.unlink()
    command = [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(pdf_path.parent), str(docx_path)]
    result = subprocess.run(command, check=False, capture_output=True, text=True, timeout=180)
    if result.returncode != 0:
        raise ValidationError(f"DOCX to PDF conversion failed: {result.stderr.strip() or result.stdout.strip()}")
    generated = pdf_path.parent / f"{docx_path.stem}.pdf"
    if generated != pdf_path and generated.exists():
        generated.replace(pdf_path)
    if not pdf_path.exists():
        raise ValidationError("DOCX to PDF conversion did not create the expected PDF")
    return {
        "source": docx_rel,
        "paired_docx": docx_rel,
        "pdf_path": pdf_rel,
        "tool": "soffice",
        "tool_path": soffice,
        "fallback_reason": reportlab_note,
        "stdout": result.stdout.strip(),
    }


def _try_external_reportlab_pdf(
    pdf_path: Path, script: dict[str, Any], primary_error: str
) -> tuple[dict[str, Any] | None, str | None]:
    root = Path(script["run_dir"])
    pdf_rel = _relative_or_absolute(root, pdf_path)
    docx_rel = project_deliverable_relpaths(root, script=script)["stage3_speaker_script_docx"]
    payload_path = root / STATE_DIR / "pdf_conversion" / "reportlab_payload.json"
    payload_path.parent.mkdir(parents=True, exist_ok=True)
    write_json(payload_path, script)

    errors: list[str] = []
    for python_path in _external_reportlab_python_candidates():
        if not _python_has_reportlab(python_path):
            errors.append(f"{python_path}: reportlab unavailable")
            continue

        helper = f"""
import json
import sys
from pathlib import Path
sys.path.insert(0, {str(Path(__file__).resolve().parents[1])!r})
from ppt_skill_v3.json_io import read_json
from ppt_skill_v3.speaker_script import _render_pdf_with_reportlab
script = read_json(Path({str(payload_path)!r}))
font_path = _render_pdf_with_reportlab(script, Path({str(pdf_path)!r}))
print(json.dumps({{"font_path": str(font_path)}}, ensure_ascii=False))
"""
        result = subprocess.run([str(python_path), "-c", helper], check=False, capture_output=True, text=True, timeout=180)
        if result.returncode == 0 and pdf_path.exists():
            try:
                details = json.loads(result.stdout.strip().splitlines()[-1])
            except Exception:
                details = {}
            return (
                {
                    "source": SCRIPT_JSON_REL,
                    "paired_docx": docx_rel,
                    "pdf_path": pdf_rel,
                    "tool": "reportlab",
                    "tool_path": str(python_path),
                    "font_path": details.get("font_path"),
                    "external_python": True,
                    "fallback_reason": f"default Python ReportLab unavailable: {primary_error}",
                },
                None,
            )
        errors.append(f"{python_path}: {result.stderr.strip() or result.stdout.strip() or 'no PDF created'}")
    return None, "; ".join(errors) if errors else "no external Python with ReportLab found"


def _external_reportlab_python_candidates() -> list[Path]:
    candidates: list[Path] = []
    env_path = os.environ.get("PPT_SKILL_PDF_PYTHON")
    if env_path:
        candidates.append(Path(env_path).expanduser())
    candidates.append(Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python3")

    unique: list[Path] = []
    seen: set[str] = set()
    current = Path(sys.executable).resolve()
    for candidate in candidates:
        if not candidate.exists():
            continue
        resolved = candidate.resolve()
        if resolved == current or str(resolved) in seen:
            continue
        seen.add(str(resolved))
        unique.append(resolved)
    return unique


def _python_has_reportlab(python_path: Path) -> bool:
    result = subprocess.run([str(python_path), "-c", "import reportlab"], check=False, capture_output=True, text=True, timeout=10)
    return result.returncode == 0


def _render_pdf_with_reportlab(script: dict[str, Any], pdf_path: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate

    font_path = _find_pdf_font_path()
    font_name = "Stage3ScriptCJK"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font_name, str(font_path)))

    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "BodyCN",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=10.8,
        leading=16.2,
        textColor=colors.HexColor("#1F2937"),
        spaceAfter=7,
        wordWrap="CJK",
    )
    title = ParagraphStyle(
        "TitleCN",
        parent=body,
        fontSize=21,
        leading=27,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#111827"),
        spaceAfter=6,
    )
    subtitle = ParagraphStyle(
        "SubtitleCN",
        parent=body,
        fontSize=10.5,
        leading=15,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#5B6978"),
        spaceAfter=6,
    )
    h1 = ParagraphStyle(
        "HeadingCN",
        parent=body,
        fontSize=14.2,
        leading=19,
        textColor=colors.HexColor("#0F4C75"),
        spaceBefore=12,
        spaceAfter=6,
        wordWrap="CJK",
    )
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(pdf_path),
        pagesize=A4,
        leftMargin=23 * mm,
        rightMargin=23 * mm,
        topMargin=22 * mm,
        bottomMargin=22 * mm,
        title=script["talk"]["title"],
        author=script["talk"]["speaker_role"],
    )

    story: list[Any] = []
    story.extend(
        [
            Paragraph("PPT 演讲逐字稿", subtitle),
            Paragraph(_p(script["talk"]["title"]), title),
            Paragraph(_p(script["talk"]["speaker_role"]), subtitle),
            HRFlowable(width="86%", thickness=1.0, color=colors.HexColor("#0F4C75"), spaceBefore=4, spaceAfter=12),
        ]
    )

    for slide in script["slides"]:
        story.append(Paragraph(_p(f"第 {slide['slide_index']:02d} 页｜{slide['slide_title']}"), h1))
        for paragraph_text in slide["script"]:
            story.append(Paragraph(_p(paragraph_text), body))

    def draw_footer(canvas, document):
        canvas.saveState()
        canvas.setFont(font_name, 8.5)
        canvas.setFillColor(colors.HexColor("#728191"))
        if document.page > 1:
            canvas.drawRightString(A4[0] - 23 * mm, A4[1] - 14 * mm, script["talk"]["title"])
            canvas.setStrokeColor(colors.HexColor("#D7E1E8"))
            canvas.line(23 * mm, A4[1] - 17 * mm, A4[0] - 23 * mm, A4[1] - 17 * mm)
        canvas.drawCentredString(A4[0] / 2, 12 * mm, f"第 {document.page} 页")
        canvas.restoreState()

    doc.build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)
    if not pdf_path.exists():
        raise ValidationError("ReportLab did not create the expected PDF")
    return font_path


def _file_entry(label: str, root: Path, relative: str) -> dict[str, str]:
    path = root / relative
    if not path.exists():
        raise FileNotFoundError(f"required speaker script file not found: {path}")
    return {
        "label": label,
        "path": relative,
        "sha256": f"sha256:{hashlib.sha256(path.read_bytes()).hexdigest()}",
    }


def _script_character_count(script: dict[str, Any]) -> int:
    chunks: list[str] = []
    for slide in script["slides"]:
        chunks.extend(slide["script"])
    return len("".join(chunks))


def _estimate_minutes(character_count: int) -> float:
    return round(max(character_count / 240, 0.1), 2)
