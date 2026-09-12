from __future__ import annotations

import hashlib
import os
import shutil
import subprocess
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .validation import ValidationError


def find_pdf_font_path() -> Path:
    env_font = os.environ.get("PPT_SKILL_CJK_FONT")
    if env_font:
        path = Path(env_font).expanduser()
        if path.exists():
            return path
    candidates = [
        Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
        Path("/Library/Fonts/Arial Unicode.ttf"),
        Path("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttf"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.otf"),
        Path("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"),
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/msyhbd.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/simsun.ttc"),
    ]
    for path in candidates:
        if path.exists():
            return path
    raise ValidationError("no embeddable CJK font found for ReportLab PDF generation")


def add_page_number(paragraph) -> None:
    paragraph.add_run("第 ")
    _add_field(paragraph, "PAGE")
    paragraph.add_run(" 页 / 共 ")
    _add_field(paragraph, "NUMPAGES")
    paragraph.add_run(" 页")


def relative_or_absolute(root: Path, path: Path) -> str:
    try:
        return str(path.relative_to(root))
    except ValueError:
        return str(path)


def pdf_page_count(path: Path) -> int | None:
    try:
        from pypdf import PdfReader

        return len(PdfReader(str(path)).pages)
    except Exception:
        pdfinfo = shutil.which("pdfinfo")
        if not pdfinfo:
            return None
        result = subprocess.run(
            [pdfinfo, str(path)],
            check=False,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=30,
        )
        if result.returncode != 0:
            return None
        for line in result.stdout.splitlines():
            if line.startswith("Pages:"):
                try:
                    return int(line.split(":", 1)[1].strip())
                except ValueError:
                    return None
        return None


def pdf_text_probe(path: Path) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        text = "\n".join((page.extract_text() or "") for page in reader.pages[:2])
        return "selectable_text_detected" if text.strip() else "no_selectable_text_detected"
    except Exception:
        pdftotext = shutil.which("pdftotext")
        if not pdftotext:
            return "not_checked"
        result = subprocess.run(
            [pdftotext, "-f", "1", "-l", "2", str(path), "-"],
            check=False,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=30,
        )
        if result.returncode != 0:
            return "not_checked"
        return "selectable_text_detected" if result.stdout.strip() else "no_selectable_text_detected"


def p(value: str) -> str:
    return escape(str(value)).replace("\n", "<br/>")


def require_locked_source_exists(root: Path, locked_source: dict[str, str]) -> None:
    source_path = resolve_locked_source_path(root, locked_source["source_path"])
    if not source_path.exists():
        raise ValidationError("locked presentation source does not exist")
    expected_sha = locked_source.get("source_sha256")
    if expected_sha:
        actual_sha = f"sha256:{hashlib.sha256(source_path.read_bytes()).hexdigest()}"
        if actual_sha != expected_sha:
            raise ValidationError("locked presentation source sha256 does not match")


def require_locked_source_matches_state(state: dict[str, Any], locked_source: dict[str, str]) -> None:
    state_source = state.get("stage3_locked_presentation_source")
    if not isinstance(state_source, dict):
        return
    for field in ("source_mode", "source_path"):
        if state_source.get(field) != locked_source.get(field):
            raise ValidationError(f"speaker_script.basis.locked_presentation_source.{field} does not match project state")
    if state_source.get("source_sha256") and state_source.get("source_sha256") != locked_source.get("source_sha256"):
        raise ValidationError("speaker_script.basis.locked_presentation_source.source_sha256 does not match project state")


def resolve_locked_source_path(root: Path, source_path: str) -> Path:
    path = Path(source_path)
    return path if path.is_absolute() else root / path


def _add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
