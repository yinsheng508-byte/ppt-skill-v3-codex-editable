from __future__ import annotations

import hashlib
import re
import shutil
import subprocess
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .deliverable_naming import project_deliverable_relpaths
from .education_context import subject_profile_for
from .events import append_event
from .json_io import read_json, write_json
from .speaker_script import (
    _add_page_number,
    _find_pdf_font_path,
    _p,
    _pdf_page_count,
    _pdf_text_probe,
    _relative_or_absolute,
    _require_locked_source_exists,
    _require_locked_source_matches_state,
)
from .state import read_state, stage4_lesson_plan_required, write_state
from .time_utils import now_iso
from .validation import (
    ValidationError,
    validate_lesson_plan,
    validate_lesson_plan_manifest,
    validate_lesson_plan_qa,
)


USER_DIR = "阶段4_演讲稿输出/教案设计"
STATE_DIR = "_state/阶段4"
LESSON_PLAN_JSON_REL = f"{STATE_DIR}/lesson_plan.json"
MANIFEST_REL = f"{STATE_DIR}/lesson_plan_manifest.json"
QA_REL = f"{STATE_DIR}/lesson_plan_qa.json"
NOTE_REL = f"{USER_DIR}/教案生成说明.md"
KEY_VALUE_TABLE_WIDTHS_CM = [3.2, 14.8]
PROCESS_TABLE_WIDTHS_CM = [2.2, 4.4, 3.6, 3.8, 4.0]
PROCESS_TABLE_TOTAL_WIDTH_CM = sum(PROCESS_TABLE_WIDTHS_CM)
FONT_FAMILY = "Arial Unicode MS"
FONT_FALLBACKS = ["Arial Unicode MS", "Noto Sans CJK SC", "Microsoft YaHei", "SimSun", "Songti SC"]
INK = RGBColor(31, 41, 55)
MUTED = RGBColor(91, 105, 120)
ACCENT = RGBColor(30, 87, 107)
ACCENT_DARK = "1E576B"
HEADER_FILL = "E8F1F4"
LIGHT_FILL = "F7FAFB"
UNKNOWN_TEXT_VALUES = {"", "待确认", "未确认", "未知", "不详", "unknown", "none", "null", "n/a", "na", "needs_confirmation"}
INTERNAL_REF_LABELS = {
    "primary_textbook": "教材相关内容",
    "teacher_reference": "教师用书相关内容",
    "exercise_reference": "练习材料",
    "sample_lesson_plan": "参考教案样例",
}
TEACHER_VISIBLE_FORBIDDEN_TERMS = (
    "Runtime",
    "Manifest",
    "QA",
    "QA草稿",
    "decision",
    "stage4_script_completed",
    "schema_version",
    "lesson_plan_qa",
    "lesson_plan_context",
    "subject_extension",
    "education_context",
)
OVERVIEW_DEPTH_RULES = {
    "textbook_analysis": {
        "label": "教材分析",
        "min_items": 3,
        "min_total_chars": 120,
        "min_item_chars": 24,
        "anchors": ("教材", "单元", "课件", "ppt", "页面", "课文", "文本", "实验", "史料", "例题", "作品", "地图", "材料", "任务"),
        "min_anchor_items": 2,
    },
    "learner_analysis": {
        "label": "学情分析",
        "min_items": 3,
        "min_total_chars": 100,
        "min_item_chars": 22,
        "anchors": ("学生", "已有", "经验", "困难", "混淆", "障碍", "误区", "支架", "课堂", "观察", "表达", "理解"),
        "min_anchor_items": 2,
    },
    "core_competency_goals": {
        "label": "核心素养目标",
        "min_items": 3,
        "min_total_chars": 120,
        "min_item_chars": 24,
        "anchors": ("能", "依据", "完成", "说明", "解释", "判断", "推理", "表达", "评价", "记录", "表格", "作品", "口头"),
        "min_anchor_items": 3,
    },
    "key_points": {
        "label": "教学重点",
        "min_items": 2,
        "min_total_chars": 56,
        "min_item_chars": 18,
        "anchors": ("知识", "方法", "任务", "活动", "成果", "文本", "实验", "史料", "例题", "材料", "概念", "结构", "语句", "题目", "含义", "现象", "结论", "纲领", "关系"),
        "min_anchor_items": 2,
    },
    "difficult_points": {
        "label": "教学难点",
        "min_items": 2,
        "min_total_chars": 56,
        "min_item_chars": 18,
        "anchors": ("学生", "能否", "容易", "困难", "混淆", "误区", "依据", "过程", "表达", "解释", "理解"),
        "min_anchor_items": 2,
    },
    "breakthrough_strategy": {
        "label": "突破策略",
        "min_items": 3,
        "min_total_chars": 100,
        "min_item_chars": 22,
        "anchors": ("问题", "任务", "学习单", "表格", "追问", "评价", "订正", "活动", "实验", "史料", "例题", "读写", "地图", "展示", "检测", "批注", "小练笔", "迁移", "记录", "板演"),
        "min_anchor_items": 3,
    },
    "teaching_methods": {
        "label": "教学方法",
        "min_items": 3,
        "min_total_chars": 12,
        "min_item_chars": 2,
        "anchors": ("问题", "任务", "史料", "实验", "探究", "合作", "展示", "评价", "检测", "朗读", "批注", "例题", "变式", "读写", "地图", "项目", "示范", "实践", "讨论"),
        "min_anchor_items": 2,
    },
}
SNAKE_CASE_RE = re.compile(r"\b[a-z][a-z0-9]+(?:_[a-z0-9]+)+\b")
INTERNAL_ID_RE = re.compile(r"\b(?:mat|tb|textbook|sample)[-_]?\d+\b", re.IGNORECASE)


def build_lesson_plan(run_dir: str | Path, lesson_plan_json: str | Path) -> dict[str, Any]:
    root = Path(run_dir)
    state = read_state(root)
    if state["current_stage"] != "stage4" or state["status"] not in {
        "ready_for_stage4_script",
        "stage4_script_generated",
        "stage4_lesson_plan_generated",
    }:
        raise ValidationError("lesson plan generation requires stage4 ready_for_stage4_script state")
    if not stage4_lesson_plan_required(state):
        raise ValidationError("lesson plan generation requires K12 stage4_outputs.lesson_plan.required=true")

    lesson_plan = validate_lesson_plan(read_json(lesson_plan_json))
    if lesson_plan["project_name"] != state["project_name"]:
        raise ValidationError("lesson_plan.project_name does not match project state")
    if lesson_plan["run_dir"] != state["run_dir"]:
        raise ValidationError("lesson_plan.run_dir does not match project state")
    if lesson_plan["education_context"].get("is_k12") is not True:
        raise ValidationError("lesson plan generation requires K12 education_context.is_k12=true")

    locked_source = lesson_plan["basis"]["locked_presentation_source"]
    _require_locked_source_exists(root, locked_source)
    _require_locked_source_matches_state(state, locked_source)

    relpaths = project_deliverable_relpaths(root, state=state, script=_topic_proxy(lesson_plan))
    markdown_rel = relpaths["stage4_lesson_plan"]
    docx_rel = relpaths["stage4_lesson_plan_docx"]
    pdf_rel = relpaths["stage4_lesson_plan_pdf"]

    (root / USER_DIR / "docx").mkdir(parents=True, exist_ok=True)
    (root / USER_DIR / "pdf").mkdir(parents=True, exist_ok=True)
    (root / STATE_DIR / "pdf_conversion").mkdir(parents=True, exist_ok=True)
    (root / STATE_DIR / "logs").mkdir(parents=True, exist_ok=True)

    markdown = _render_markdown(lesson_plan)
    _assert_teacher_visible_clean(markdown)
    write_json(root / LESSON_PLAN_JSON_REL, lesson_plan)
    (root / markdown_rel).write_text(markdown, encoding="utf-8")
    _render_docx(lesson_plan, root / docx_rel)
    conversion = _build_pdf(root / docx_rel, root / pdf_rel, lesson_plan)
    conversion["pdf_path"] = pdf_rel
    if conversion.get("paired_docx"):
        conversion["paired_docx"] = docx_rel
    if conversion.get("source") in {str(root / LESSON_PLAN_JSON_REL), _relative_or_absolute(root, root / LESSON_PLAN_JSON_REL)}:
        conversion["source"] = LESSON_PLAN_JSON_REL

    profile = subject_profile_for(lesson_plan["education_context"])
    layout_probe = _docx_layout_probe(lesson_plan)
    summary = {
        "periods": len(lesson_plan["periods"]),
        "activities": sum(len(period["process"]) for period in lesson_plan["periods"]),
        "subject_group": profile["subject_group"],
        "word_table_mode": "a4_portrait_compact_process_tables",
        "layout_probe": layout_probe,
        "overview_depth_probe": _overview_depth_probe(lesson_plan),
        "pdf_text_probe": _pdf_text_probe(root / pdf_rel),
        "docx_font_family": FONT_FAMILY,
        "docx_font_fallbacks": FONT_FALLBACKS,
        "docx_language": "zh-CN",
        "locked_source_mode": locked_source["source_mode"],
        "locked_source_path": locked_source["source_path"],
    }
    pdf_pages = _pdf_page_count(root / pdf_rel)
    if pdf_pages is not None:
        summary["pdf_pages"] = pdf_pages
    manifest = validate_lesson_plan_manifest(
        {
            "schema_version": "1.0",
            "project_name": state["project_name"],
            "run_dir": state["run_dir"],
            "files": [
                _file_entry("Markdown 教案设计", root, markdown_rel),
                _file_entry("Word 教案设计", root, docx_rel),
                _file_entry("PDF 教案设计", root, pdf_rel),
            ],
            "conversion": conversion,
            "summary": summary,
            "created_at": now_iso(),
            "status": "generated",
        }
    )
    write_json(root / MANIFEST_REL, manifest)
    qa = _build_runtime_qa(state, lesson_plan, manifest, profile)
    write_json(root / QA_REL, qa)
    _write_generation_note(root, lesson_plan, manifest, qa, locked_source)

    state["status"] = "stage4_lesson_plan_generated"
    state["required_actor"] = "main_controller"
    stage4_outputs = state.setdefault("stage4_outputs", {})
    lesson_plan_output = stage4_outputs.setdefault("lesson_plan", {})
    lesson_plan_output["required"] = True
    lesson_plan_output["status"] = "generated"
    state["stage4_lesson_plan_required"] = True
    state.setdefault("user_artifacts", {})["stage4_lesson_plan"] = markdown_rel
    state["user_artifacts"]["stage4_lesson_plan_docx"] = docx_rel
    state["user_artifacts"]["stage4_lesson_plan_pdf"] = pdf_rel
    state.setdefault("expected_user_paths", {})["stage4_lesson_plan"] = markdown_rel
    state["expected_user_paths"]["stage4_lesson_plan_docx"] = docx_rel
    state["expected_user_paths"]["stage4_lesson_plan_pdf"] = pdf_rel
    state.setdefault("quality", {})["stage4"] = "pending_controller_review"
    state["next_required_action"] = "主控大模型检查阶段4教案设计、Word 和 PDF；如讲稿尚未生成，继续生成讲稿；全部通过后记录阶段4输出完成决策"
    write_state(root, state)
    append_event(root, "lesson_plan_generated", "runtime", periods=len(lesson_plan["periods"]))
    return manifest


def _topic_proxy(lesson_plan: dict[str, Any]) -> dict[str, Any]:
    return {"talk": {"title": _display_lesson_title(lesson_plan)}}


def _docx_layout_probe(lesson_plan: dict[str, Any]) -> dict[str, Any]:
    process_lengths = [len(period["process"]) for period in lesson_plan["periods"]]
    return {
        "page_size": "A4",
        "page_orientation": "portrait",
        "process_section_orientation": "portrait",
        "process_table_split_mode": "natural_page_split",
        "process_table_count": len(process_lengths),
        "process_table_repeat_header": True,
        "process_table_start_policy": "natural_flow_after_teaching_preparation",
        "manual_process_table_chunks": 0,
        "max_steps_per_process_table": max(process_lengths) if process_lengths else 0,
        "max_process_steps_in_period": max(process_lengths) if process_lengths else 0,
        "process_table_width_cm": round(PROCESS_TABLE_TOTAL_WIDTH_CM, 2),
        "process_table_column_widths_cm": PROCESS_TABLE_WIDTHS_CM,
        "font_fallbacks": FONT_FALLBACKS,
    }


def _lesson_title(lesson_plan: dict[str, Any]) -> str:
    context = lesson_plan.get("education_context", {})
    title = context.get("lesson_title") if isinstance(context, dict) else None
    return str(title).strip() if isinstance(title, str) and title.strip() else lesson_plan["project_name"]


def _display_lesson_title(lesson_plan: dict[str, Any]) -> str:
    context = lesson_plan.get("education_context", {})
    raw_title = _lesson_title(lesson_plan)
    if not isinstance(context, dict):
        return _normalize_lesson_title_spacing(raw_title)
    topic = _lesson_topic_title(context) or raw_title
    lesson_number = _lesson_number_label(context)
    if not lesson_number:
        return _normalize_lesson_title_spacing(topic)
    if _title_starts_with_lesson_number(raw_title, lesson_number):
        return _normalize_lesson_title_spacing(raw_title)
    return _normalize_lesson_title_spacing(f"{lesson_number} {topic}")


def _lesson_topic_title(context: dict[str, Any]) -> str:
    title = _text(context.get("lesson_title", ""))
    if not _is_visible_value(title):
        return ""
    lesson_number = _lesson_number_label(context)
    if lesson_number:
        title = _strip_lesson_number_prefix(title, lesson_number)
    return _normalize_lesson_title_spacing(title)


def _lesson_number_label(context: dict[str, Any]) -> str:
    lesson_number = _text(context.get("lesson_number", ""))
    if not _is_visible_value(lesson_number):
        return ""
    compact = re.sub(r"\s+", "", lesson_number)
    if re.fullmatch(r"第.+课", compact):
        return compact
    if re.fullmatch(r"[0-9一二三四五六七八九十]+", compact):
        return f"第{compact}课"
    return lesson_number.strip()


def _title_starts_with_lesson_number(title: str, lesson_number: str) -> bool:
    if not title or not lesson_number:
        return False
    return _strip_lesson_number_prefix(title, lesson_number) != _normalize_lesson_title_spacing(title)


def _strip_lesson_number_prefix(title: str, lesson_number: str) -> str:
    normalized = _normalize_lesson_title_spacing(title)
    patterns = [rf"^{re.escape(lesson_number)}\s*"]
    if lesson_number.startswith("第") and lesson_number.endswith("课"):
        middle = re.escape(lesson_number[1:-1])
        patterns.append(rf"^第\s*{middle}\s*课\s*")
    for pattern in patterns:
        stripped = re.sub(pattern, "", normalized).strip()
        if stripped != normalized:
            return _normalize_lesson_title_spacing(stripped)
    return normalized


def _normalize_lesson_title_spacing(value: Any) -> str:
    text = str(value).strip()
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"^第\s*([0-9一二三四五六七八九十]+)\s*课\s*", r"第\1课 ", text)
    return text.strip()


def _period_heading(period: dict[str, Any], lesson_plan: dict[str, Any]) -> str:
    title = _text(period.get("period_title", ""))
    total_periods = len(lesson_plan.get("periods", []))
    if total_periods == 1:
        return f"本课教学安排｜{title}" if title else "本课教学安排"
    prefix = f"第{period['period_no']}课时"
    return f"{prefix}｜{title}" if title else prefix


def _render_markdown(lesson_plan: dict[str, Any]) -> str:
    context = lesson_plan["education_context"]
    overview = lesson_plan["overview"]
    lines = [
        f"# {_display_lesson_title(lesson_plan)}｜教案设计",
        "",
        "## 基本信息",
        "",
        "| 项目 | 内容 |",
        "|---|---|",
    ]
    for key, value in _basic_info_rows(context):
        lines.append(f"| {_md_cell(key)} | {_md_cell(value)} |")
    lines.extend(["", "## 教学分析", ""])
    for title, field in [
        ("教材分析", "textbook_analysis"),
        ("学情分析", "learner_analysis"),
        ("核心素养目标", "core_competency_goals"),
        ("教学重点", "key_points"),
        ("教学难点", "difficult_points"),
        ("突破策略", "breakthrough_strategy"),
        ("教学方法", "teaching_methods"),
    ]:
        if overview.get(field):
            lines.extend([f"### {title}", ""])
            lines.extend([f"- {_md_cell(item)}" for item in _as_text_list(overview[field])])
            lines.append("")
    preparation_rows = _preparation_rows(overview.get("preparation", {}))
    if preparation_rows:
        lines.extend(["### 教学准备", "", "| 类型 | 内容 |", "|---|---|"])
        for key, value in preparation_rows:
            lines.append(f"| {_md_cell(key)} | {_md_cell(value)} |")
        lines.append("")

    periods = lesson_plan["periods"]
    single_period = len(periods) == 1
    for period in periods:
        section_prefix = "##" if single_period else "###"
        if not single_period:
            lines.extend(
                [
                    f"## {_period_heading(period, lesson_plan)}",
                    "",
                    f"- 建议时长：{period['duration_minutes']}分钟",
                ]
            )
            if period.get("slide_range"):
                lines.append(f"- 课件范围：{period['slide_range']}")
            if period.get("material_scope"):
                material_scope = _join_visible_refs(period["material_scope"])
                if material_scope:
                    lines.append(f"- 材料范围：{material_scope}")
            if period.get("activity_overview"):
                lines.append(f"- 活动概览：{period['activity_overview']}")
            lines.extend(["", "### 课时目标", ""])
            lines.extend([f"- {_md_cell(item)}" for item in period["period_objectives"]])
        lines.extend(["", f"{section_prefix} 教学过程", ""])
        lines.extend(
            [
                "| 环节 | 教师活动 | 学生活动 | 任务与评价 | 设计意图/二次备课 |",
                "|---|---|---|---|---|",
            ]
        )
        for step in period["process"]:
            lines.append(
                "| "
                + " | ".join(
                    [
                        _md_cell(step["phase"]),
                        _md_cell(step["teacher_activity"]),
                        _md_cell(step["student_activity"]),
                        _md_cell(_learning_assessment_text(step)),
                        _md_cell(_intent_note_text(step)),
                    ]
                )
                + " |"
            )
        lines.extend(["", f"{section_prefix} 作业设计", "", "| 类型 | 内容 |", "|---|---|"])
        for key, value in _homework_rows(period["homework"]):
            lines.append(f"| {_md_cell(key)} | {_md_cell(value)} |")
        if period.get("blackboard_design"):
            lines.extend(["", f"{section_prefix} 板书设计", ""])
            lines.extend([f"- {_md_cell(item)}" for item in period["blackboard_design"]])
        lines.extend(["", f"{section_prefix} 教学反思", ""])
        for key, value in _reflection_rows(period["reflection"]):
            lines.append(f"- **{_md_cell(key)}**：{_md_cell(value)}")
        lines.append("")

    return "\n".join(lines).rstrip() + "\n"


def _render_docx(lesson_plan: dict[str, Any], target: Path) -> None:
    doc = Document()
    _configure_document(doc, lesson_plan)
    _configure_styles(doc)
    _write_docx_header(doc, lesson_plan)
    _write_basic_info(doc, lesson_plan)
    _write_overview(doc, lesson_plan)
    _write_periods(doc, lesson_plan)
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)


def _configure_document(doc: Document, lesson_plan: dict[str, Any]) -> None:
    section = doc.sections[0]
    _configure_section_page(section, landscape=False)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)
    section.different_first_page_header_footer = True

    header = section.header.paragraphs[0]
    header.text = f"{_display_lesson_title(lesson_plan)}｜教案设计"
    header.style = doc.styles["Header"]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_page_number(footer)


def _configure_section_page(section, *, landscape: bool) -> None:
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21)
        section.top_margin = Cm(1.4)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.4)
        section.right_margin = Cm(1.4)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)


def _configure_styles(doc: Document) -> None:
    styles = doc.styles
    _set_style_font(styles["Normal"], FONT_FAMILY, 10.5, INK)
    styles["Normal"].paragraph_format.line_spacing = 1.35
    styles["Normal"].paragraph_format.space_after = Pt(4)

    for name, size, color in [
        ("Title", 22, RGBColor(17, 24, 39)),
        ("Subtitle", 10.5, MUTED),
        ("Heading 1", 15, ACCENT),
        ("Heading 2", 12, RGBColor(17, 24, 39)),
        ("Heading 3", 10.5, ACCENT),
    ]:
        _set_style_font(styles[name], FONT_FAMILY, size, color, bold=name != "Subtitle")
        styles[name].paragraph_format.keep_with_next = True
    styles["Title"].paragraph_format.space_after = Pt(6)
    styles["Subtitle"].paragraph_format.space_after = Pt(6)
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(6)
    styles["Heading 2"].paragraph_format.space_before = Pt(8)
    styles["Heading 2"].paragraph_format.space_after = Pt(4)
    styles["Heading 3"].paragraph_format.space_before = Pt(6)
    styles["Heading 3"].paragraph_format.space_after = Pt(3)

    for style_name, size, color in [
        ("Lesson Body", 10.5, INK),
        ("Lesson Table", 8.5, INK),
        ("Lesson Muted", 9.0, MUTED),
    ]:
        style = _paragraph_style(doc, style_name)
        _set_style_font(style, FONT_FAMILY, size, color)
        style.paragraph_format.line_spacing = 1.2
        style.paragraph_format.space_after = Pt(2)


def _write_docx_header(doc: Document, lesson_plan: dict[str, Any]) -> None:
    context = lesson_plan["education_context"]
    label = doc.add_paragraph("课堂教学设计", style="Subtitle")
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title = doc.add_paragraph(_display_lesson_title(lesson_plan), style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = doc.add_paragraph(_meta_line(context), style="Lesson Muted")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_rule(doc, ACCENT_DARK)


def _write_basic_info(doc: Document, lesson_plan: dict[str, Any]) -> None:
    doc.add_heading("基本信息", level=1)
    _add_key_value_table(doc, _basic_info_rows(lesson_plan["education_context"]))


def _write_overview(doc: Document, lesson_plan: dict[str, Any]) -> None:
    overview = lesson_plan["overview"]
    doc.add_heading("教学分析与目标", level=1)
    for title, field in [
        ("教材分析", "textbook_analysis"),
        ("学情分析", "learner_analysis"),
        ("核心素养目标", "core_competency_goals"),
        ("教学重点", "key_points"),
        ("教学难点", "difficult_points"),
        ("突破策略", "breakthrough_strategy"),
        ("教学方法", "teaching_methods"),
    ]:
        if overview.get(field):
            doc.add_heading(title, level=2)
            _add_bullets(doc, _as_text_list(overview[field]))
    preparation_rows = _preparation_rows(overview.get("preparation", {}))
    if preparation_rows:
        doc.add_heading("教学准备", level=2)
        _add_key_value_table(doc, preparation_rows)


def _write_periods(doc: Document, lesson_plan: dict[str, Any]) -> None:
    periods = lesson_plan["periods"]
    single_period = len(periods) == 1
    for period in periods:
        section_level = 1 if single_period else 2
        if not single_period:
            doc.add_heading(_period_heading(period, lesson_plan), level=1)
            _add_key_value_table(doc, _period_info_rows(period))
            doc.add_heading("课时目标", level=2)
            _add_bullets(doc, period["period_objectives"])
        doc.add_heading("教学过程", level=section_level)
        _add_process_table(doc, period["process"])
        doc.add_heading("作业设计", level=section_level)
        _add_key_value_table(doc, _homework_rows(period["homework"]))
        if period.get("blackboard_design"):
            doc.add_heading("板书设计", level=section_level)
            _add_bullets(doc, period["blackboard_design"])
        doc.add_heading("教学反思", level=section_level)
        _add_reflection_block(doc, _reflection_rows(period["reflection"]))


def _add_key_value_table(doc: Document, rows: list[tuple[str, Any]]) -> None:
    rows = _visible_rows(rows)
    if not rows:
        return
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    _set_table_fixed_widths(table, KEY_VALUE_TABLE_WIDTHS_CM)
    _set_repeat_table_header(table.rows[0])
    _set_cell(table.rows[0].cells[0], "项目", bold=True, fill=HEADER_FILL)
    _set_cell(table.rows[0].cells[1], "内容", bold=True, fill=HEADER_FILL)
    for key, value in rows:
        cells = table.add_row().cells
        _set_cell(cells[0], key, bold=True, fill=LIGHT_FILL)
        _set_cell(cells[1], _text(value), size=9.5)
    _set_table_fixed_widths(table, KEY_VALUE_TABLE_WIDTHS_CM)


def _add_process_table(doc: Document, steps: list[dict[str, Any]]) -> None:
    headers = ["环节", "教师活动", "学生活动", "任务与评价", "设计意图/二次备课"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    _set_table_fixed_widths(table, PROCESS_TABLE_WIDTHS_CM)
    _set_repeat_table_header(table.rows[0])
    for cell, header in zip(table.rows[0].cells, headers):
        _set_cell(cell, header, bold=True, fill=HEADER_FILL, size=8.0, compact=True)
    for step in steps:
        row = table.add_row()
        values = [
            step["phase"],
            step["teacher_activity"],
            step["student_activity"],
            _learning_assessment_text(step),
            _intent_note_text(step),
        ]
        for cell, value in zip(row.cells, values):
            _set_cell(cell, value, size=7.4, compact=True)
    _set_table_fixed_widths(table, PROCESS_TABLE_WIDTHS_CM)


def _add_reflection_block(doc: Document, rows: list[tuple[str, Any]]) -> None:
    for key, value in _visible_rows(rows):
        paragraph = doc.add_paragraph(style="Lesson Body")
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(3)
        label_run = paragraph.add_run(f"{_text(key)}：")
        _set_run_style(label_run, bold=True, size=10.0, color=ACCENT)
        values = _as_text_list(value)
        if not values:
            continue
        first_run = paragraph.add_run(values[0])
        _set_run_style(first_run, size=10.0, color=INK)
        for item in values[1:]:
            paragraph.add_run("\n")
            item_run = paragraph.add_run(item)
            _set_run_style(item_run, size=10.0, color=INK)


def _set_cell(cell, value: Any, *, bold: bool = False, fill: str | None = None, size: float = 9.0, compact: bool = False) -> None:
    cell.text = ""
    if fill:
        _shade_cell(cell, fill)
    _set_cell_margins(cell, compact=compact)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    paragraph = cell.paragraphs[0]
    paragraph.style = "Lesson Table"
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0 if compact else 1)
    paragraph.paragraph_format.line_spacing = 1.05 if compact else 1.15
    run = paragraph.add_run(_text(value) or " ")
    run.bold = bold
    run.font.name = FONT_FAMILY
    run.font.size = Pt(size)
    run.font.color.rgb = INK
    _set_run_east_asia(run, FONT_FAMILY)


def _set_run_style(run, *, bold: bool = False, size: float = 10.5, color: RGBColor = INK) -> None:
    run.bold = bold
    run.font.name = FONT_FAMILY
    run.font.size = Pt(size)
    run.font.color.rgb = color
    _set_run_east_asia(run, FONT_FAMILY)


def _set_table_fixed_widths(table, widths_cm: list[float]) -> None:
    table.autofit = False
    _set_table_width(table, sum(widths_cm))
    _set_table_grid(table, widths_cm)
    for row in table.rows:
        for index, width_cm in enumerate(widths_cm):
            if index >= len(row.cells):
                break
            cell = row.cells[index]
            cell.width = Cm(width_cm)
            _set_cell_width(cell, width_cm)


def _set_table_width(table, width_cm: float) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        _insert_child_before(
            tbl_pr,
            tbl_w,
            [
                "w:jc",
                "w:tblCellSpacing",
                "w:tblInd",
                "w:tblBorders",
                "w:shd",
                "w:tblLayout",
                "w:tblCellMar",
                "w:tblLook",
            ],
        )
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(_cm_to_twips(width_cm)))


def _set_table_grid(table, widths_cm: list[float]) -> None:
    tbl = table._tbl
    tbl_grid = tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        tbl.insert(1, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width_cm in widths_cm:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(_cm_to_twips(width_cm)))
        tbl_grid.append(grid_col)


def _set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.insert(0, tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(_cm_to_twips(width_cm)))


def _set_cell_margins(cell, *, compact: bool = False) -> None:
    margin_pt = 1.2 if compact else 2.2
    margin_twips = _pt_to_twips(margin_pt)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        _insert_child_before(tc_pr, tc_mar, ["w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"])
    for side in ["top", "left", "bottom", "right"]:
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_twips))
        node.set(qn("w:type"), "dxa")


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tbl_header = OxmlElement("w:tblHeader")
        _insert_child_before(tr_pr, tbl_header, ["w:tblCellSpacing", "w:jc", "w:hidden"])


def _insert_child_before(parent, child, before_tags: list[str]) -> None:
    targets = {qn(tag) for tag in before_tags}
    for index, existing in enumerate(parent):
        if existing.tag in targets:
            parent.insert(index, child)
            return
    parent.append(child)


def _cm_to_twips(value: float) -> int:
    return int(round(value * 567))


def _pt_to_twips(value: float) -> int:
    return int(round(value * 20))


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        insert_before = {"w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"}
        for index, child in enumerate(tc_pr):
            if child.tag in {qn(tag) for tag in insert_before}:
                tc_pr.insert(index, shd)
                break
        else:
            tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)


def _add_bullets(doc: Document, values: list[Any]) -> None:
    for value in _as_text_list(values):
        paragraph = doc.add_paragraph(str(value), style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        for run in paragraph.runs:
            run.font.name = FONT_FAMILY
            run.font.size = Pt(10.5)
            run.font.color.rgb = INK
            _set_run_east_asia(run, FONT_FAMILY)


def _paragraph_style(doc: Document, name: str):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def _set_style_font(style, font_name: str, size_pt: float, color: RGBColor, *, bold: bool = False) -> None:
    style.font.name = font_name
    style.font.size = Pt(size_pt)
    style.font.color.rgb = color
    style.font.bold = bold
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:eastAsia"), font_name)
    r_fonts.set(qn("w:cs"), font_name)
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")


def _set_run_east_asia(run, font_name: str) -> None:
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), font_name)


def _add_rule(doc: Document, color: str) -> None:
    paragraph = doc.add_paragraph(style="Lesson Muted")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run("-" * 64)
    run.font.name = FONT_FAMILY
    run.font.size = Pt(5)
    run.font.color.rgb = RGBColor.from_string(color)
    _set_run_east_asia(run, FONT_FAMILY)


def _build_pdf(docx_path: Path, pdf_path: Path, lesson_plan: dict[str, Any]) -> dict[str, Any]:
    root = Path(lesson_plan["run_dir"])
    docx_rel = _relative_or_absolute(root, docx_path)
    pdf_rel = _relative_or_absolute(root, pdf_path)
    try:
        font_path = _render_pdf_with_reportlab(lesson_plan, pdf_path)
        return {
            "source": LESSON_PLAN_JSON_REL,
            "paired_docx": docx_rel,
            "pdf_path": pdf_rel,
            "tool": "reportlab",
            "font_path": str(font_path),
        }
    except Exception as reportlab_error:
        reportlab_note = str(reportlab_error)

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise ValidationError(f"ReportLab PDF generation failed and LibreOffice/soffice is not available: {reportlab_note}")
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    if pdf_path.exists():
        pdf_path.unlink()
    command = [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(pdf_path.parent), str(docx_path)]
    result = subprocess.run(command, check=False, capture_output=True, text=True, timeout=180)
    if result.returncode != 0:
        raise ValidationError(f"DOCX to PDF conversion failed: {result.stderr.strip() or result.stdout.strip()}")
    generated = pdf_path.parent / f"{docx_path.stem}.pdf"
    if generated != pdf_path and generated.exists():
        generated.replace(pdf_path)
    if not pdf_path.exists():
        raise ValidationError("DOCX to PDF conversion did not create the expected PDF")
    return {
        "source": docx_rel,
        "paired_docx": docx_rel,
        "pdf_path": pdf_rel,
        "tool": "soffice",
        "tool_path": soffice,
        "fallback_reason": reportlab_note,
        "stdout": result.stdout.strip(),
    }


def _render_pdf_with_reportlab(lesson_plan: dict[str, Any], pdf_path: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import HRFlowable, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font_path = _find_pdf_font_path()
    font_name = "Stage4LessonCJK"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font_name, str(font_path)))

    styles = getSampleStyleSheet()
    body = ParagraphStyle("BodyCN", parent=styles["BodyText"], fontName=font_name, fontSize=9.4, leading=13.5, textColor=colors.HexColor("#1F2937"), wordWrap="CJK")
    small = ParagraphStyle("SmallCN", parent=body, fontSize=8.0, leading=10.5)
    title = ParagraphStyle("TitleCN", parent=body, fontSize=21, leading=27, alignment=TA_CENTER, textColor=colors.HexColor("#111827"), spaceAfter=6)
    subtitle = ParagraphStyle("SubtitleCN", parent=body, fontSize=9.5, leading=13, alignment=TA_CENTER, textColor=colors.HexColor("#5B6978"), spaceAfter=5)
    h1 = ParagraphStyle("H1CN", parent=body, fontSize=13.5, leading=18, textColor=colors.HexColor("#1E576B"), spaceBefore=10, spaceAfter=5)
    h2 = ParagraphStyle("H2CN", parent=body, fontSize=10.5, leading=14.5, textColor=colors.HexColor("#111827"), spaceBefore=6, spaceAfter=3)

    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(pdf_path),
        pagesize=A4,
        leftMargin=17 * mm,
        rightMargin=17 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=_display_lesson_title(lesson_plan),
        author="ppt-skill-v2",
    )
    story: list[Any] = [
        Paragraph("课堂教学设计", subtitle),
        Paragraph(_p(_display_lesson_title(lesson_plan)), title),
        Paragraph(_p(_meta_line(lesson_plan["education_context"])), subtitle),
        HRFlowable(width="88%", thickness=1.0, color=colors.HexColor("#1E576B"), spaceBefore=3, spaceAfter=10),
        Paragraph("基本信息", h1),
        _pdf_key_value_table(_basic_info_rows(lesson_plan["education_context"]), body, small),
        Paragraph("教学分析与目标", h1),
    ]
    overview = lesson_plan["overview"]
    for section_title, field in [
        ("教材分析", "textbook_analysis"),
        ("学情分析", "learner_analysis"),
        ("核心素养目标", "core_competency_goals"),
        ("教学重点", "key_points"),
        ("教学难点", "difficult_points"),
        ("突破策略", "breakthrough_strategy"),
        ("教学方法", "teaching_methods"),
    ]:
        if overview.get(field):
            story.append(Paragraph(section_title, h2))
            story.extend(Paragraph(_p(f"• {item}"), body) for item in _as_text_list(overview[field]))
    preparation_rows = _preparation_rows(overview.get("preparation", {}))
    if preparation_rows:
        story.append(Paragraph("教学准备", h2))
        story.append(_pdf_key_value_table(preparation_rows, body, small))

    periods = lesson_plan["periods"]
    single_period = len(periods) == 1
    for period in periods:
        section_style = h1 if single_period else h2
        if not single_period:
            story.append(Paragraph(_p(_period_heading(period, lesson_plan)), h1))
            story.append(_pdf_key_value_table(_period_info_rows(period), body, small))
            story.append(Paragraph("课时目标", h2))
            story.extend(Paragraph(_p(f"• {item}"), body) for item in period["period_objectives"])
        story.append(KeepTogether([Paragraph("教学过程", section_style), _pdf_process_table(period["process"], small)]))
        story.append(KeepTogether([Paragraph("作业设计", section_style), _pdf_key_value_table(_homework_rows(period["homework"]), body, small)]))
        if period.get("blackboard_design"):
            story.append(Paragraph("板书设计", section_style))
            story.extend(Paragraph(_p(f"• {item}"), body) for item in period["blackboard_design"])
        reflection_flow = _pdf_reflection_block(_reflection_rows(period["reflection"]), body)
        if reflection_flow:
            story.append(KeepTogether([Paragraph("教学反思", section_style), reflection_flow[0]]))
            story.extend(reflection_flow[1:])
        story.append(Spacer(1, 4))

    def draw_footer(canvas, document):
        canvas.saveState()
        canvas.setFont(font_name, 8)
        canvas.setFillColor(colors.HexColor("#728191"))
        if document.page > 1:
            canvas.drawRightString(A4[0] - 17 * mm, A4[1] - 11 * mm, _display_lesson_title(lesson_plan))
            canvas.setStrokeColor(colors.HexColor("#D7E1E8"))
            canvas.line(17 * mm, A4[1] - 14 * mm, A4[0] - 17 * mm, A4[1] - 14 * mm)
        canvas.drawCentredString(A4[0] / 2, 10 * mm, f"第 {document.page} 页")
        canvas.restoreState()

    doc.build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)
    if not pdf_path.exists():
        raise ValidationError("ReportLab did not create the expected PDF")
    return font_path


def _pdf_key_value_table(rows: list[tuple[str, Any]], body_style, small_style):
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Table, TableStyle

    rows = _visible_rows(rows)
    data = [[Paragraph("<b>项目</b>", small_style), Paragraph("<b>内容</b>", small_style)]]
    data.extend([[_pdf_cell(key, small_style), _pdf_cell(value, body_style)] for key, value in rows])
    table = Table(data, colWidths=[29 * mm, 132 * mm], repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B9CAD1")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8F1F4")),
                ("BACKGROUND", (0, 1), (0, -1), colors.HexColor("#F7FAFB")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def _pdf_process_table(steps: list[dict[str, Any]], small_style):
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Table, TableStyle

    headers = ["环节", "教师活动", "学生活动", "任务与评价", "设计意图/二次备课"]
    data = [[Paragraph(f"<b>{_p(header)}</b>", small_style) for header in headers]]
    for step in steps:
        data.append(
            [
                _pdf_cell(step["phase"], small_style),
                _pdf_cell(step["teacher_activity"], small_style),
                _pdf_cell(step["student_activity"], small_style),
                _pdf_cell(_learning_assessment_text(step), small_style),
                _pdf_cell(_intent_note_text(step), small_style),
            ]
        )
    table = Table(data, colWidths=[20 * mm, 40 * mm, 34 * mm, 35 * mm, 32 * mm], repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B9CAD1")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8F1F4")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    return table


def _pdf_reflection_block(rows: list[tuple[str, Any]], body_style):
    from reportlab.platypus import Paragraph, Spacer

    flowables: list[Any] = []
    for key, value in _visible_rows(rows):
        values = _as_text_list(value)
        if not values:
            continue
        lines = [f"<b>{_p(key)}：</b>{_p(values[0])}"]
        lines.extend(f"• {_p(item)}" for item in values[1:])
        flowables.append(Paragraph("<br/>".join(lines), body_style))
        flowables.append(Spacer(1, 1))
    return flowables


def _pdf_cell(value: Any, style):
    from reportlab.platypus import Paragraph

    return Paragraph(_p(_text(value) or " "), style)


def _build_runtime_qa(
    state: dict[str, Any], lesson_plan: dict[str, Any], manifest: dict[str, Any], profile: dict[str, Any]
) -> dict[str, Any]:
    teacher_items = _teacher_confirmation_items(lesson_plan, profile)
    source_conflicts = lesson_plan["basis"].get("source_conflicts", [])
    warnings: list[str] = []
    if profile.get("warning"):
        warnings.append(str(profile["warning"]))
    if manifest["summary"].get("pdf_text_probe") == "no_selectable_text_detected":
        warnings.append("教案 PDF 暂未检测到可选择文本，完成前需主控复核转换结果。")
    return validate_lesson_plan_qa(
        {
            "schema_version": "1.0",
            "project_name": state["project_name"],
            "run_dir": state["run_dir"],
            "status": "needs_controller_review",
            "blockers": [],
            "warnings": warnings,
            "teacher_confirmation_items": teacher_items,
            "source_conflicts": source_conflicts,
            "review_notes": ["运行时已生成教案三件套；主控需按 K12 教案规范完成语义自审。"],
            "controller_reviewed": False,
            "created_at": now_iso(),
        }
    )


def _teacher_confirmation_items(lesson_plan: dict[str, Any], profile: dict[str, Any]) -> list[str]:
    context = lesson_plan["education_context"]
    items: list[str] = []
    for field, label in [
        ("grade", "年级"),
        ("subject", "学科"),
        ("textbook_version", "教材版本"),
        ("lesson_title", "课题"),
        ("period_count", "课时数"),
    ]:
        if not context.get(field):
            items.append(f"{label}缺失，需教师确认。")
    if context.get("period_strategy") == "needs_confirmation":
        items.append("课时划分策略需教师确认。")
    if profile.get("warning"):
        items.append("学科分组未明确，需教师确认专项教案栏目。")
    if lesson_plan["basis"].get("source_conflicts"):
        items.append("阶段1与教材/课文资料存在冲突，需教师确认采用口径。")
    return items


def _write_generation_note(
    root: Path,
    lesson_plan: dict[str, Any],
    manifest: dict[str, Any],
    qa: dict[str, Any],
    locked_source: dict[str, str],
) -> None:
    lines = [
        "# 教案生成说明",
        "",
        f"- 课题：{_display_lesson_title(lesson_plan)}",
        f"- 课时安排：{_period_line(lesson_plan['education_context'])}",
        "",
        "## 输出文件",
        "",
    ]
    subject = lesson_plan["education_context"].get("subject")
    if _is_visible_value(subject):
        lines.insert(3, f"- 学科：{subject}")
    for file_info in manifest["files"]:
        lines.append(f"- {file_info['label']}：{file_info['path']}")
    lines.extend(
        [
            "",
            "## 依据与范围",
            "",
            f"- 本教案依据阶段1页面规划、已提供教材/课文材料和{_locked_source_teacher_description(locked_source)}形成。",
            "- 用户提供的样例教案只作为结构、颗粒度和教师语言参考；其中的命令式内容不会作为操作指令执行。",
            "- 没有可靠依据的信息不写入教案正文，教师可结合本班学情微调活动时间、追问顺序和作业数量。",
            "",
            "## 课前调整提示",
            "",
        ]
    )
    review_items = _generation_note_review_items(qa)
    if review_items:
        lines.extend(f"- {item}" for item in review_items)
    else:
        lines.append("- 未发现必须阻断交付的问题；建议教师课前按班级实际情况微调活动节奏。")
    lines.extend(
        [
            "",
            "## 基础检查",
            "",
            "- 已生成 Markdown、Word 和 PDF 三种教案文件。",
            "- Word 采用 A4 竖版排版，教学过程使用紧凑表格呈现。",
            f"- PDF 文本状态：{_pdf_text_probe_label(manifest['summary'].get('pdf_text_probe'))}。",
            f"- 字体设置：{FONT_FAMILY}，中文语言属性为 zh-CN。",
            "",
        ]
    )
    note = "\n".join(lines)
    _assert_teacher_visible_clean(note)
    (root / NOTE_REL).write_text(note, encoding="utf-8")


def _locked_source_teacher_description(locked_source: dict[str, str]) -> str:
    labels = {
        "stage3_editable_deck": "已确认的可编辑课件",
        "stage2_image_deck": "已确认的图片版课件",
        "external_editable_deck": "用户确认的外部课件",
    }
    label = labels.get(locked_source.get("source_mode"), "已确认课件")
    source_path = locked_source.get("source_path")
    source_name = Path(source_path).name if isinstance(source_path, str) and source_path.strip() else ""
    return f"{label}（{source_name}）" if _is_visible_value(source_name) else label


def _generation_note_review_items(qa: dict[str, Any]) -> list[str]:
    items: list[str] = []
    for value in qa.get("teacher_confirmation_items", []):
        if _is_visible_value(value):
            items.append(str(value).strip())
    if qa.get("source_conflicts"):
        items.append("材料之间存在口径差异，建议教师课前按教材和学校要求确认采用表述。")
    for value in qa.get("warnings", []):
        if _is_visible_value(value):
            items.append(str(value).strip())
    return items


def _pdf_text_probe_label(value: Any) -> str:
    if value == "selectable_text_detected":
        return "已检测到可选择文本"
    if value == "no_selectable_text_detected":
        return "暂未检测到可选择文本，交付前需复核"
    return "已生成，建议打开抽查"


def _basic_info_rows(context: dict[str, Any]) -> list[tuple[str, Any]]:
    return _visible_rows([
        ("年级", context.get("grade", "")),
        ("学科", context.get("subject", "")),
        ("学段", _school_stage_label(context.get("school_stage"))),
        ("教材版本", context.get("textbook_version", "")),
        ("单元", context.get("unit", "")),
        ("课题", _lesson_topic_title(context) or context.get("lesson_title", "")),
        ("教材课序", _lesson_number_label(context)),
        ("课型", context.get("class_type", "")),
        ("课时安排", _period_line(context)),
    ])


def _preparation_rows(preparation: dict[str, Any]) -> list[tuple[str, Any]]:
    return _visible_rows([
        ("教师准备", preparation.get("teacher", [])),
        ("学生准备", preparation.get("student", [])),
        ("资源材料", preparation.get("resources", [])),
        ("设备器材", preparation.get("equipment", [])),
        ("安全提示", preparation.get("safety_notes", [])),
    ])


def _homework_rows(homework: dict[str, Any]) -> list[tuple[str, Any]]:
    return _visible_rows([
        ("基础作业", homework.get("basic", [])),
        ("巩固作业", homework.get("consolidation", [])),
        ("拓展作业", homework.get("extension", [])),
        ("参考答案/评价要点", homework.get("answer_key", [])),
    ])


def _reflection_rows(reflection: dict[str, Any]) -> list[tuple[str, Any]]:
    note = (
        "以下为课后填写提示，供教师结合真实课堂表现记录，不预设实际教学结果。"
        if reflection.get("mode") == "pre_teaching_prompt"
        else "请基于真实课堂实施情况记录有效做法、学生困难和后续调整。"
    )
    return _visible_rows([
        ("填写说明", note),
        ("课堂观察重点", reflection.get("success_observation_prompts", [])),
        ("可能问题", reflection.get("risk_observation_prompts", [])),
        ("调整策略", reflection.get("improvement_prompts", [])),
        (
            "二次备课记录方向",
            reflection.get("secondary_preparation_prompts", [])
            or reflection.get("followup_prompts", []),
        ),
    ])


def _period_info_rows(period: dict[str, Any]) -> list[tuple[str, Any]]:
    return _visible_rows(
        [
            ("建议时长", f"{period['duration_minutes']}分钟"),
            ("课件范围", period.get("slide_range", "")),
            ("材料范围", _join_visible_refs(period.get("material_scope", []))),
            ("活动概览", period.get("activity_overview", "")),
        ]
    )


def _period_line(context: dict[str, Any]) -> str:
    period_count = context.get("period_count")
    minutes = context.get("minutes_per_period")
    if period_count and minutes:
        try:
            period_total = int(period_count)
        except (TypeError, ValueError):
            period_total = None
        if period_total == 1:
            return f"本课共{period_count}课时，建议{minutes}分钟完成"
        return f"本课共{period_count}课时，建议每课时{minutes}分钟"
    if period_count:
        return f"本课共{period_count}课时"
    return ""


def _meta_line(context: dict[str, Any]) -> str:
    parts = [
        context.get("grade"),
        context.get("subject"),
        context.get("textbook_version"),
        context.get("unit"),
        _period_line(context),
    ]
    return "｜".join(str(part).strip() for part in parts if _is_visible_value(part))


def _refs_text(step: dict[str, Any]) -> str:
    parts: list[str] = []
    if step.get("ppt_slide_refs"):
        parts.append("PPT：" + "、".join(str(item) for item in step["ppt_slide_refs"]))
    if step.get("materials_refs"):
        refs = _join_visible_refs(step["materials_refs"])
        if refs:
            parts.append("材料：" + refs)
    return "\n".join(parts)


def _join_visible_refs(value: Any) -> str:
    refs: list[str] = []
    for item in _as_text_list(value):
        ref = _visible_material_ref(item)
        if ref:
            refs.append(ref)
    return "；".join(refs)


def _visible_material_ref(value: Any) -> str:
    text = _text(value)
    if not text:
        return ""
    lowered = text.strip().lower()
    if lowered in INTERNAL_REF_LABELS:
        return INTERNAL_REF_LABELS[lowered]
    with_label = re.match(r"^\s*(?:mat|tb|textbook|sample)[-_]?\d+\s*[:：]\s*(.+)$", text, flags=re.IGNORECASE)
    if with_label:
        return with_label.group(1).strip()
    if INTERNAL_ID_RE.fullmatch(text.strip()):
        return "教材相关内容"
    if SNAKE_CASE_RE.fullmatch(text.strip()):
        return ""
    return text


def _learning_assessment_text(step: dict[str, Any]) -> str:
    parts: list[str] = []
    if step.get("learning_task"):
        parts.append("学习任务：" + _join_list(step["learning_task"]))
    if step.get("assessment_evidence"):
        parts.append("评价证据：" + _join_list(step["assessment_evidence"]))
    if step.get("safety_notes"):
        parts.append("安全：" + _join_list(step["safety_notes"]))
    return "\n".join(parts)


def _intent_note_text(step: dict[str, Any]) -> str:
    parts = [step["design_intent"]]
    if step.get("secondary_preparation_note"):
        parts.append("二次备课：" + step["secondary_preparation_note"])
    return "\n".join(parts)


def _as_text_list(value: Any) -> list[str]:
    if isinstance(value, list):
        return [_text(item) for item in value if _is_visible_value(_text(item))]
    if isinstance(value, dict):
        return [f"{key}：{_text(child)}" for key, child in value.items()]
    text = _text(value)
    return [text] if _is_visible_value(text) else []


def _join_list(value: Any) -> str:
    return "；".join(_as_text_list(value))


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "\n".join(_as_text_list(value))
    if isinstance(value, dict):
        return "\n".join(f"{key}：{_text(child)}" for key, child in value.items())
    return str(value).strip()


def _md_cell(value: Any) -> str:
    text = _text(value)
    return text.replace("|", "\\|").replace("\n", "<br>")


def _school_stage_label(value: Any) -> str:
    return {
        "primary": "小学",
        "junior_high": "初中",
        "senior_high": "高中",
        "unknown": "",
    }.get(value, str(value).strip() if value else "")


def _visible_rows(rows: list[tuple[str, Any]]) -> list[tuple[str, Any]]:
    return [(key, value) for key, value in rows if _is_visible_value(value)]


def _is_visible_value(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, list):
        return any(_is_visible_value(item) for item in value)
    if isinstance(value, dict):
        return any(_is_visible_value(item) for item in value.values())
    text = str(value).strip()
    return text.lower() not in UNKNOWN_TEXT_VALUES


def _overview_depth_probe(lesson_plan: dict[str, Any]) -> dict[str, Any]:
    overview = lesson_plan.get("overview", {}) if isinstance(lesson_plan.get("overview"), dict) else {}
    fields: dict[str, Any] = {}
    warnings: list[str] = []
    for field, rule in OVERVIEW_DEPTH_RULES.items():
        values = _as_text_list(overview.get(field, []))
        char_counts = [_visible_char_count(value) for value in values]
        anchor_items = _specific_anchor_item_count(values, rule.get("anchors", ()))
        field_warnings: list[str] = []
        if len(values) < rule["min_items"]:
            field_warnings.append("内容条目偏少，建议主控复核是否需要展开。")
        if sum(char_counts) < rule["min_total_chars"]:
            field_warnings.append("整体说明偏短，建议主控复核是否能支撑教师备课。")
        if any(count < rule["min_item_chars"] for count in char_counts):
            field_warnings.append("存在过短表述，建议主控复核是否过于提纲化。")
        if anchor_items < rule.get("min_anchor_items", 0):
            field_warnings.append("具体课堂依据或活动指向偏弱，建议主控复核。")
        if field_warnings:
            warnings.append(f"{rule['label']}：{'；'.join(field_warnings)}")
        fields[field] = {
            "label": rule["label"],
            "items": len(values),
            "total_chars": sum(char_counts),
            "min_item_chars": min(char_counts) if char_counts else 0,
            "specific_anchor_items": anchor_items,
            "advisory_items": rule["min_items"],
            "advisory_total_chars": rule["min_total_chars"],
            "advisory_anchor_items": rule.get("min_anchor_items", 0),
            "warnings": field_warnings,
        }
    return {"status": "pass" if not warnings else "needs_controller_review", "fields": fields, "warnings": warnings}


def _visible_char_count(value: Any) -> int:
    text = re.sub(r"\s+", "", _text(value))
    return len(text)


def _specific_anchor_item_count(values: list[str], anchors: Any) -> int:
    anchor_list = [str(anchor).lower() for anchor in anchors if str(anchor).strip()]
    if not anchor_list:
        return 0
    count = 0
    for value in values:
        text = _text(value).lower()
        if any(anchor in text for anchor in anchor_list):
            count += 1
    return count


def _teacher_visible_artifacts(markdown: str) -> list[str]:
    artifacts: list[str] = []
    for token in UNKNOWN_TEXT_VALUES:
        if _contains_teacher_visible_token(markdown, token):
            artifacts.append(token)
    for token in TEACHER_VISIBLE_FORBIDDEN_TERMS:
        if _contains_teacher_visible_token(markdown, token) and token not in artifacts:
            artifacts.append(token)
    for match in SNAKE_CASE_RE.finditer(markdown):
        token = match.group(0)
        if token not in artifacts:
            artifacts.append(token)
    for match in INTERNAL_ID_RE.finditer(markdown):
        token = match.group(0)
        if token not in artifacts:
            artifacts.append(token)
    return artifacts


def _contains_teacher_visible_token(text: str, token: str) -> bool:
    if not token:
        return False
    if re.fullmatch(r"[A-Za-z0-9_/-]+", token):
        return bool(re.search(rf"(?<![A-Za-z0-9_/-]){re.escape(token)}(?![A-Za-z0-9_/-])", text, re.IGNORECASE))
    return token in text


def _assert_teacher_visible_clean(markdown: str) -> None:
    artifacts = _teacher_visible_artifacts(markdown)
    if artifacts:
        preview = "、".join(artifacts[:8])
        raise ValidationError(f"teacher-visible lesson plan contains internal or uncertain artifacts: {preview}")


def _file_entry(label: str, root: Path, relative: str) -> dict[str, str]:
    path = root / relative
    if not path.exists():
        raise FileNotFoundError(f"required lesson plan file not found: {path}")
    return {
        "label": label,
        "path": relative,
        "sha256": f"sha256:{hashlib.sha256(path.read_bytes()).hexdigest()}",
    }
