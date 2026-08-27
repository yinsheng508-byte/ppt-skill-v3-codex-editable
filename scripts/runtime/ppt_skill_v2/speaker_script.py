from __future__ import annotations

import hashlib
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .deliverable_naming import (
    LEGACY_STAGE4_DOCX_REL,
    LEGACY_STAGE4_MARKDOWN_REL,
    LEGACY_STAGE4_PDF_REL,
    project_deliverable_relpaths,
)
from .events import append_event
from .json_io import read_json, write_json
from .state import read_state, write_state
from .time_utils import now_iso
from .validation import ValidationError, validate_speaker_script, validate_speaker_script_manifest


USER_DIR = "阶段4_演讲稿输出"
STATE_DIR = "_state/阶段4"
MARKDOWN_REL = LEGACY_STAGE4_MARKDOWN_REL
DOCX_REL = LEGACY_STAGE4_DOCX_REL
PDF_REL = LEGACY_STAGE4_PDF_REL
NOTE_REL = f"{USER_DIR}/讲稿生成说明.md"
SCRIPT_JSON_REL = f"{STATE_DIR}/speaker_script.json"
MANIFEST_REL = f"{STATE_DIR}/speaker_script_manifest.json"
FONT_FAMILY = "Arial Unicode MS"
INK = RGBColor(31, 41, 55)
MUTED = RGBColor(91, 105, 120)
ACCENT = RGBColor(15, 76, 117)
ACCENT_DARK = "0F4C75"


def build_speaker_script(run_dir: str | Path, script_json: str | Path) -> dict[str, Any]:
    root = Path(run_dir)
    state = read_state(root)
    if state["current_stage"] != "stage4" or state["status"] not in {"ready_for_stage4_script", "stage4_script_generated"}:
        raise ValidationError("speaker script generation requires stage4 ready_for_stage4_script state")

    script = validate_speaker_script(read_json(script_json))
    if script["project_name"] != state["project_name"]:
        raise ValidationError("speaker_script.project_name does not match project state")
    if script["run_dir"] != state["run_dir"]:
        raise ValidationError("speaker_script.run_dir does not match project state")
    locked_source = _locked_presentation_source(script)
    _require_locked_source_exists(root, locked_source)
    _require_locked_source_matches_state(state, locked_source)

    relpaths = project_deliverable_relpaths(root, state=state, script=script)
    markdown_rel = relpaths["stage4_speaker_script"]
    docx_rel = relpaths["stage4_speaker_script_docx"]
    pdf_rel = relpaths["stage4_speaker_script_pdf"]

    (root / USER_DIR / "docx").mkdir(parents=True, exist_ok=True)
    (root / USER_DIR / "pdf").mkdir(parents=True, exist_ok=True)
    (root / STATE_DIR / "pdf_conversion").mkdir(parents=True, exist_ok=True)
    (root / STATE_DIR / "logs").mkdir(parents=True, exist_ok=True)

    write_json(root / SCRIPT_JSON_REL, script)
    markdown = _render_markdown(script)
    (root / markdown_rel).write_text(markdown, encoding="utf-8")
    _render_docx(script, root / docx_rel)
    conversion = _build_pdf(root / docx_rel, root / pdf_rel, script)
    conversion["pdf_path"] = pdf_rel
    if conversion.get("paired_docx"):
        conversion["paired_docx"] = docx_rel
    if conversion.get("source") in {DOCX_REL, str(root / DOCX_REL)}:
        conversion["source"] = docx_rel
    script_character_count = _script_character_count(script)

    manifest = validate_speaker_script_manifest(
        {
            "schema_version": "2.0",
            "project_name": state["project_name"],
            "run_dir": state["run_dir"],
            "files": [
                _file_entry("Markdown 逐字稿", root, markdown_rel),
                _file_entry("Word 逐字稿", root, docx_rel),
                _file_entry("PDF 逐字稿", root, pdf_rel),
            ],
            "conversion": conversion,
            "summary": {
                "slides": len(script["slides"]),
                "script_characters": script_character_count,
                "estimated_minutes": _estimate_minutes(script_character_count),
                "pdf_pages": _pdf_page_count(root / pdf_rel),
                "pdf_text_probe": _pdf_text_probe(root / pdf_rel),
                "first_page_header_mode": "inline_top_header_no_cover_page",
                "docx_font_family": FONT_FAMILY,
                "docx_language": "zh-CN",
                "locked_source_mode": locked_source["source_mode"],
                "locked_source_path": locked_source["source_path"],
            },
            "created_at": now_iso(),
            "status": "generated",
        }
    )
    write_json(root / MANIFEST_REL, manifest)
    _write_generation_note(root, script, manifest, locked_source)

    state["status"] = "stage4_script_generated"
    state["required_actor"] = "main_controller"
    state.setdefault("user_artifacts", {})["stage4_speaker_script"] = markdown_rel
    state["user_artifacts"]["stage4_speaker_script_docx"] = docx_rel
    state["user_artifacts"]["stage4_speaker_script_pdf"] = pdf_rel
    state.setdefault("expected_user_paths", {})["stage4_speaker_script"] = markdown_rel
    state["expected_user_paths"]["stage4_speaker_script_docx"] = docx_rel
    state["expected_user_paths"]["stage4_speaker_script_pdf"] = pdf_rel
    state["quality"]["stage4"] = "pending_controller_review"
    state["next_required_action"] = "主控大模型检查阶段4讲稿、Word 和 PDF；通过后记录 stage4_script_completed 决策"
    write_state(root, state)
    append_event(root, "speaker_script_generated", "runtime", slides=len(script["slides"]))
    return manifest


def _locked_presentation_source(script: dict[str, Any]) -> dict[str, str]:
    basis = script["basis"]
    source = basis.get("locked_presentation_source")
    if isinstance(source, dict):
        result = {
            "source_mode": str(source["source_mode"]),
            "source_path": str(source["source_path"]),
        }
        if source.get("source_sha256"):
            result["source_sha256"] = str(source["source_sha256"])
        if source.get("confirmation_basis"):
            result["confirmation_basis"] = str(source["confirmation_basis"])
        return result
    result = {
        "source_mode": "stage3_editable_deck",
        "source_path": str(basis["stage3_editable_deck"]),
        "confirmation_basis": "用户确认阶段3可编辑 PPT，可作为阶段4讲稿锁定稿。",
    }
    if basis.get("stage3_editable_deck_sha256"):
        result["source_sha256"] = str(basis["stage3_editable_deck_sha256"])
    return result


def _require_locked_source_exists(root: Path, locked_source: dict[str, str]) -> None:
    source_path = _resolve_locked_source_path(root, locked_source["source_path"])
    if not source_path.exists():
        raise ValidationError("locked presentation source does not exist")
    expected_sha = locked_source.get("source_sha256")
    if expected_sha:
        actual_sha = f"sha256:{hashlib.sha256(source_path.read_bytes()).hexdigest()}"
        if actual_sha != expected_sha:
            raise ValidationError("locked presentation source sha256 does not match")


def _require_locked_source_matches_state(state: dict[str, Any], locked_source: dict[str, str]) -> None:
    state_source = state.get("stage4_locked_presentation_source")
    if not isinstance(state_source, dict):
        return
    for field in ("source_mode", "source_path"):
        if state_source.get(field) != locked_source.get(field):
            raise ValidationError(f"speaker_script.basis.locked_presentation_source.{field} does not match project state")
    if state_source.get("source_sha256") and state_source.get("source_sha256") != locked_source.get("source_sha256"):
        raise ValidationError("speaker_script.basis.locked_presentation_source.source_sha256 does not match project state")


def _resolve_locked_source_path(root: Path, source_path: str) -> Path:
    path = Path(source_path)
    return path if path.is_absolute() else root / path


def _render_markdown(script: dict[str, Any]) -> str:
    talk = script["talk"]
    lines = [
        f"# {talk['title']}演讲逐字稿",
        "",
        f"讲者身份：{talk['speaker_role']}",
        "",
        "---",
        "",
    ]
    for slide in script["slides"]:
        lines.extend(
            [
                f"### 第 {slide['slide_index']:02d} 页｜{slide['slide_title']}",
                "",
            ]
        )
        lines.extend([f"{paragraph}" for paragraph in slide["script"]])
        lines.append("")
    lines.append("")
    return "\n".join(lines)


def _render_docx(script: dict[str, Any], target: Path) -> None:
    doc = Document()
    _configure_document(doc, script)
    _configure_styles(doc)
    _write_document_header(doc, script)
    _write_slide_scripts(doc, script)
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)


def _configure_document(doc: Document, script: dict[str, Any]) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.1)
    section.different_first_page_header_footer = True

    header = section.header.paragraphs[0]
    header.text = script["talk"]["title"]
    header.style = doc.styles["Header"]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_page_number(footer)


def _configure_styles(doc: Document) -> None:
    styles = doc.styles
    _set_style_font(styles["Normal"], FONT_FAMILY, 10.5, INK)
    styles["Normal"].paragraph_format.line_spacing = 1.4
    styles["Normal"].paragraph_format.space_after = Pt(4)

    for name, size, color in [
        ("Title", 22, RGBColor(17, 24, 39)),
        ("Subtitle", 12, MUTED),
        ("Heading 1", 15, ACCENT),
        ("Heading 2", 11, RGBColor(17, 24, 39)),
    ]:
        _set_style_font(styles[name], FONT_FAMILY, size, color, bold=name in {"Title", "Heading 1", "Heading 2"})
        styles[name].paragraph_format.keep_with_next = True
    styles["Title"].paragraph_format.space_before = Pt(0)
    styles["Title"].paragraph_format.space_after = Pt(5)
    styles["Subtitle"].paragraph_format.space_before = Pt(0)
    styles["Subtitle"].paragraph_format.space_after = Pt(4)
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(7)
    styles["Heading 2"].paragraph_format.space_before = Pt(8)
    styles["Heading 2"].paragraph_format.space_after = Pt(4)

    script_body = _paragraph_style(doc, "Script Body")
    _set_style_font(script_body, FONT_FAMILY, 11, INK)
    script_body.paragraph_format.line_spacing = 1.45
    script_body.paragraph_format.space_after = Pt(7)

    note_style = _paragraph_style(doc, "Speaker Note")
    _set_style_font(note_style, FONT_FAMILY, 9.5, MUTED)
    note_style.paragraph_format.line_spacing = 1.25
    note_style.paragraph_format.space_after = Pt(4)

    meta_style = _paragraph_style(doc, "Talk Meta")
    _set_style_font(meta_style, FONT_FAMILY, 9.5, MUTED)
    meta_style.paragraph_format.space_after = Pt(5)


def _write_document_header(doc: Document, script: dict[str, Any]) -> None:
    talk = script["talk"]
    label = doc.add_paragraph("PPT 演讲逐字稿", style="Subtitle")
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title = doc.add_paragraph(talk["title"], style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph(style="Talk Meta")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(talk["speaker_role"])
    meta.paragraph_format.space_after = Pt(6)
    _add_rule(doc, ACCENT_DARK)


def _write_slide_scripts(doc: Document, script: dict[str, Any]) -> None:
    for slide in script["slides"]:
        doc.add_heading(f"第 {slide['slide_index']:02d} 页｜{slide['slide_title']}", level=1)
        for paragraph_text in slide["script"]:
            doc.add_paragraph(paragraph_text, style="Script Body")


def _paragraph_style(doc: Document, name: str):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def _set_style_font(style, font_name: str, size_pt: float, color: RGBColor, *, bold: bool = False) -> None:
    style.font.name = font_name
    style.font.size = Pt(size_pt)
    style.font.color.rgb = color
    style.font.bold = bold
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:eastAsia"), font_name)
    r_fonts.set(qn("w:cs"), font_name)
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")


def _paragraph_border(paragraph, *, left: str | None = None, bottom: str | None = None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    for side, color in {"left": left, "bottom": bottom}.items():
        if not color:
            continue
        border = p_bdr.find(qn(f"w:{side}"))
        if border is None:
            border = OxmlElement(f"w:{side}")
            p_bdr.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8" if side == "left" else "6")
        border.set(qn("w:space"), "6" if side == "left" else "1")
        border.set(qn("w:color"), color)


def _add_rule(doc: Document, color: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(10)
    _paragraph_border(paragraph, bottom=color)


def _find_pdf_font_path() -> Path:
    candidates = [
        Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
        Path("/Library/Fonts/Arial Unicode.ttf"),
        Path("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttf"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.otf"),
        Path("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"),
    ]
    for path in candidates:
        if path.exists():
            return path
    raise ValidationError("no embeddable CJK font found for ReportLab PDF generation")


def _add_page_number(paragraph) -> None:
    paragraph.add_run("第 ")
    _add_field(paragraph, "PAGE")
    paragraph.add_run(" 页 / 共 ")
    _add_field(paragraph, "NUMPAGES")
    paragraph.add_run(" 页")


def _add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def _build_pdf(docx_path: Path, pdf_path: Path, script: dict[str, Any]) -> dict[str, Any]:
    root = Path(script["run_dir"])
    docx_rel = _relative_or_absolute(root, docx_path)
    pdf_rel = _relative_or_absolute(root, pdf_path)
    try:
        font_path = _render_pdf_with_reportlab(script, pdf_path)
        return {
            "source": SCRIPT_JSON_REL,
            "paired_docx": docx_rel,
            "pdf_path": pdf_rel,
            "tool": "reportlab",
            "font_path": str(font_path),
        }
    except Exception as reportlab_error:
        reportlab_note = str(reportlab_error)

    external_conversion, external_error = _try_external_reportlab_pdf(pdf_path, script, reportlab_note)
    if external_conversion:
        return external_conversion
    if external_error:
        reportlab_note = f"{reportlab_note}; external ReportLab attempt failed: {external_error}"

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise ValidationError(f"ReportLab PDF generation failed and LibreOffice/soffice is not available: {reportlab_note}")
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    if pdf_path.exists():
        pdf_path.unlink()
    command = [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(pdf_path.parent), str(docx_path)]
    result = subprocess.run(command, check=False, capture_output=True, text=True, timeout=180)
    if result.returncode != 0:
        raise ValidationError(f"DOCX to PDF conversion failed: {result.stderr.strip() or result.stdout.strip()}")
    generated = pdf_path.parent / f"{docx_path.stem}.pdf"
    if generated != pdf_path and generated.exists():
        generated.replace(pdf_path)
    if not pdf_path.exists():
        raise ValidationError("DOCX to PDF conversion did not create the expected PDF")
    return {
        "source": docx_rel,
        "paired_docx": docx_rel,
        "pdf_path": pdf_rel,
        "tool": "soffice",
        "tool_path": soffice,
        "fallback_reason": reportlab_note,
        "stdout": result.stdout.strip(),
    }


def _try_external_reportlab_pdf(
    pdf_path: Path, script: dict[str, Any], primary_error: str
) -> tuple[dict[str, Any] | None, str | None]:
    root = Path(script["run_dir"])
    pdf_rel = _relative_or_absolute(root, pdf_path)
    docx_rel = project_deliverable_relpaths(root, script=script)["stage4_speaker_script_docx"]
    payload_path = root / STATE_DIR / "pdf_conversion" / "reportlab_payload.json"
    payload_path.parent.mkdir(parents=True, exist_ok=True)
    write_json(payload_path, script)

    errors: list[str] = []
    for python_path in _external_reportlab_python_candidates():
        if not _python_has_reportlab(python_path):
            errors.append(f"{python_path}: reportlab unavailable")
            continue

        helper = f"""
import json
import sys
from pathlib import Path
sys.path.insert(0, {str(Path(__file__).resolve().parents[1])!r})
from ppt_skill_v2.json_io import read_json
from ppt_skill_v2.speaker_script import _render_pdf_with_reportlab
script = read_json(Path({str(payload_path)!r}))
font_path = _render_pdf_with_reportlab(script, Path({str(pdf_path)!r}))
print(json.dumps({{"font_path": str(font_path)}}, ensure_ascii=False))
"""
        result = subprocess.run([str(python_path), "-c", helper], check=False, capture_output=True, text=True, timeout=180)
        if result.returncode == 0 and pdf_path.exists():
            try:
                details = json.loads(result.stdout.strip().splitlines()[-1])
            except Exception:
                details = {}
            return (
                {
                    "source": SCRIPT_JSON_REL,
                    "paired_docx": docx_rel,
                    "pdf_path": pdf_rel,
                    "tool": "reportlab",
                    "tool_path": str(python_path),
                    "font_path": details.get("font_path"),
                    "external_python": True,
                    "fallback_reason": f"default Python ReportLab unavailable: {primary_error}",
                },
                None,
            )
        errors.append(f"{python_path}: {result.stderr.strip() or result.stdout.strip() or 'no PDF created'}")
    return None, "; ".join(errors) if errors else "no external Python with ReportLab found"


def _external_reportlab_python_candidates() -> list[Path]:
    candidates: list[Path] = []
    env_path = os.environ.get("PPT_SKILL_PDF_PYTHON")
    if env_path:
        candidates.append(Path(env_path).expanduser())
    candidates.append(Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python3")

    unique: list[Path] = []
    seen: set[str] = set()
    current = Path(sys.executable).resolve()
    for candidate in candidates:
        if not candidate.exists():
            continue
        resolved = candidate.resolve()
        if resolved == current or str(resolved) in seen:
            continue
        seen.add(str(resolved))
        unique.append(resolved)
    return unique


def _python_has_reportlab(python_path: Path) -> bool:
    result = subprocess.run([str(python_path), "-c", "import reportlab"], check=False, capture_output=True, text=True, timeout=10)
    return result.returncode == 0


def _render_pdf_with_reportlab(script: dict[str, Any], pdf_path: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate

    font_path = _find_pdf_font_path()
    font_name = "Stage4CJK"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font_name, str(font_path)))

    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "BodyCN",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=10.8,
        leading=16.2,
        textColor=colors.HexColor("#1F2937"),
        spaceAfter=7,
        wordWrap="CJK",
    )
    title = ParagraphStyle(
        "TitleCN",
        parent=body,
        fontSize=21,
        leading=27,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#111827"),
        spaceAfter=6,
    )
    subtitle = ParagraphStyle(
        "SubtitleCN",
        parent=body,
        fontSize=10.5,
        leading=15,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#5B6978"),
        spaceAfter=6,
    )
    h1 = ParagraphStyle(
        "HeadingCN",
        parent=body,
        fontSize=14.2,
        leading=19,
        textColor=colors.HexColor("#0F4C75"),
        spaceBefore=12,
        spaceAfter=6,
        wordWrap="CJK",
    )
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(pdf_path),
        pagesize=A4,
        leftMargin=23 * mm,
        rightMargin=23 * mm,
        topMargin=22 * mm,
        bottomMargin=22 * mm,
        title=script["talk"]["title"],
        author=script["talk"]["speaker_role"],
    )

    story: list[Any] = []
    story.extend(
        [
            Paragraph("PPT 演讲逐字稿", subtitle),
            Paragraph(_p(script["talk"]["title"]), title),
            Paragraph(_p(script["talk"]["speaker_role"]), subtitle),
            HRFlowable(width="86%", thickness=1.0, color=colors.HexColor("#0F4C75"), spaceBefore=4, spaceAfter=12),
        ]
    )

    for slide in script["slides"]:
        story.append(Paragraph(_p(f"第 {slide['slide_index']:02d} 页｜{slide['slide_title']}"), h1))
        for paragraph_text in slide["script"]:
            story.append(Paragraph(_p(paragraph_text), body))

    def draw_footer(canvas, document):
        canvas.saveState()
        canvas.setFont(font_name, 8.5)
        canvas.setFillColor(colors.HexColor("#728191"))
        if document.page > 1:
            canvas.drawRightString(A4[0] - 23 * mm, A4[1] - 14 * mm, script["talk"]["title"])
            canvas.setStrokeColor(colors.HexColor("#D7E1E8"))
            canvas.line(23 * mm, A4[1] - 17 * mm, A4[0] - 23 * mm, A4[1] - 17 * mm)
        canvas.drawCentredString(A4[0] / 2, 12 * mm, f"第 {document.page} 页")
        canvas.restoreState()

    doc.build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)
    if not pdf_path.exists():
        raise ValidationError("ReportLab did not create the expected PDF")
    return font_path


def _write_generation_note(root: Path, script: dict[str, Any], manifest: dict[str, Any], locked_source: dict[str, str]) -> None:
    lines = [
        "# 讲稿生成说明",
        "",
        f"- 演讲标题：{script['talk']['title']}",
        f"- 讲者身份：{script['talk']['speaker_role']}",
        f"- 逐字稿页数：{manifest['summary']['slides']} 页",
        "",
        "## 输出文件",
        "",
    ]
    for file_info in manifest["files"]:
        lines.append(f"- {file_info['label']}：{file_info['path']}")
    lines.extend(
        [
            "",
            "## 生成路线",
            "",
            f"- 主控大模型基于已确认锁定稿撰写逐页演讲稿：{_locked_source_description(locked_source)}。",
            "- 逐字稿按 PPT 页码自然展开；主控按内容需要判断是否补充权威资料，不改写 PPT 结论。",
            "- Runtime 基于 `speaker_script.json` 生成 Markdown、Word 和 PDF。",
            "- Word/PDF 不设置单独封面页；演讲标题和讲者身份作为第一页顶部题头，随后直接进入逐页正文。",
            f"- Word 使用结构化样式、中文字体与语言属性（{FONT_FAMILY} / zh-CN）。",
            "- PDF 文字可选择，不使用整页图片拼接；默认嵌入中文字体，必要时回退为 Word 转换。",
            "",
            "## 检查状态",
            "",
            f"- Manifest：{MANIFEST_REL}",
            "- 主控检查通过后记录 `stage4_script_completed` 决策，项目完成。",
            "",
        ]
    )
    (root / NOTE_REL).write_text("\n".join(lines), encoding="utf-8")


def _locked_source_description(locked_source: dict[str, str]) -> str:
    labels = {
        "stage3_editable_deck": "阶段3可编辑 PPT",
        "stage2_image_deck": "阶段2图片版 PDF（跳过 Skill 阶段3）",
        "external_editable_deck": "外部工具生成并由用户确认的锁定稿",
    }
    label = labels.get(locked_source["source_mode"], locked_source["source_mode"])
    return f"{label}（{locked_source['source_path']}）"


def _file_entry(label: str, root: Path, relative: str) -> dict[str, str]:
    path = root / relative
    if not path.exists():
        raise FileNotFoundError(f"required speaker script file not found: {path}")
    return {
        "label": label,
        "path": relative,
        "sha256": f"sha256:{hashlib.sha256(path.read_bytes()).hexdigest()}",
    }


def _relative_or_absolute(root: Path, path: Path) -> str:
    try:
        return str(path.relative_to(root))
    except ValueError:
        return str(path)


def _script_character_count(script: dict[str, Any]) -> int:
    chunks: list[str] = []
    for slide in script["slides"]:
        chunks.extend(slide["script"])
    return len("".join(chunks))


def _pdf_page_count(path: Path) -> int | None:
    try:
        from pypdf import PdfReader

        return len(PdfReader(str(path)).pages)
    except Exception:
        pdfinfo = shutil.which("pdfinfo")
        if not pdfinfo:
            return None
        result = subprocess.run([pdfinfo, str(path)], check=False, capture_output=True, text=True, timeout=30)
        if result.returncode != 0:
            return None
        for line in result.stdout.splitlines():
            if line.startswith("Pages:"):
                try:
                    return int(line.split(":", 1)[1].strip())
                except ValueError:
                    return None
        return None


def _pdf_text_probe(path: Path) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        text = "\n".join((page.extract_text() or "") for page in reader.pages[:2])
        return "selectable_text_detected" if text.strip() else "no_selectable_text_detected"
    except Exception:
        pdftotext = shutil.which("pdftotext")
        if not pdftotext:
            return "not_checked"
        result = subprocess.run([pdftotext, "-f", "1", "-l", "2", str(path), "-"], check=False, capture_output=True, text=True, timeout=30)
        if result.returncode != 0:
            return "not_checked"
        return "selectable_text_detected" if result.stdout.strip() else "no_selectable_text_detected"


def _p(value: str) -> str:
    return escape(str(value)).replace("\n", "<br/>")


def _estimate_minutes(character_count: int) -> float:
    return round(max(character_count / 240, 0.1), 2)
