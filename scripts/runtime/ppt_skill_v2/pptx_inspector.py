from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

from pptx import Presentation

from .editable_coordinate_plan import validate_coordinate_execution_report, validate_editable_coordinate_plan
from .json_io import read_json, write_json


def build_coordinate_execution_report(
    pptx_file: str | Path,
    coordinate_plan_file: str | Path,
    output_file: str | Path | None = None,
) -> Path | dict[str, Any]:
    pptx_path = Path(pptx_file)
    plan_path = Path(coordinate_plan_file)
    plan = validate_editable_coordinate_plan(read_json(plan_path))
    prs = Presentation(str(pptx_path))
    report = {
        "schema_version": "2.0",
        "actual_source": "pptx_inspect",
        "pptx_sha256": _sha256_file(pptx_path),
        "slides": [],
    }
    for plan_slide in plan["slides"]:
        slide_index = plan_slide["slide_index"]
        pptx_slide = prs.slides[slide_index - 1]
        slide_width = int(prs.slide_width)
        slide_height = int(prs.slide_height)
        used_shape_ids: set[int] = set()
        unit_reports = []
        for text_unit in plan_slide["text_units"]:
            shape = _match_coordinate_text_shape(pptx_slide, text_unit, used_shape_ids, slide_width, slide_height)
            unit_reports.append(_inspect_coordinate_text_unit(shape, text_unit, slide_width, slide_height))
        merged_candidates = _possible_merged_text_shapes(pptx_slide, plan_slide["text_units"])
        report["slides"].append(
            {
                "slide_index": slide_index,
                "coordinate_text_units": unit_reports,
                "summary": _coordinate_slide_summary(plan_slide, unit_reports, merged_candidates),
                "possible_merged_text_shapes": merged_candidates,
            }
        )
    validate_coordinate_execution_report(report, plan)
    if output_file is None:
        return report
    output_path = Path(output_file)
    write_json(output_path, report)
    return output_path


def _match_coordinate_text_shape(
    slide: Any,
    text_unit: dict[str, Any],
    used_shape_ids: set[int],
    slide_width: int,
    slide_height: int,
) -> Any | None:
    shape = _match_shape_by_name(slide, text_unit["text_unit_id"])
    if shape is not None:
        used_shape_ids.add(id(shape))
        return shape
    candidates = []
    for candidate in slide.shapes:
        if id(candidate) in used_shape_ids or not getattr(candidate, "has_text_frame", False):
            continue
        if _shape_text(candidate).strip() == text_unit["text"]:
            candidates.append(candidate)
    if not candidates:
        return None
    planned_box = text_unit["relative_box"]
    best = min(
        candidates,
        key=lambda candidate: _relative_box_distance(
            planned_box,
            _shape_relative_box(candidate, slide_width, slide_height),
        ),
    )
    used_shape_ids.add(id(best))
    return best


def _match_shape_by_name(slide: Any, name: str) -> Any | None:
    for shape in slide.shapes:
        if getattr(shape, "name", "") == name:
            return shape
    return None


def _inspect_coordinate_text_unit(shape: Any | None, planned: dict[str, Any], slide_width: int, slide_height: int) -> dict[str, Any]:
    planned_report = _planned_coordinate_text_report(planned)
    if shape is None:
        return {
            "text_unit_id": planned["text_unit_id"],
            "planned": planned_report,
            "actual": {},
            "within_allowed_range": False,
            "deviation_reason": "coordinate text unit not found in PPTX inspect",
        }
    actual = {
        "relative_box": _shape_relative_box(shape, slide_width, slide_height),
        "font": _shape_font(shape),
    }
    within, reason = _text_within_plan(planned, actual)
    report = {
        "text_unit_id": planned["text_unit_id"],
        "planned": planned_report,
        "actual": actual,
        "within_allowed_range": within,
    }
    if not within:
        report["deviation_reason"] = reason
    return report


def _planned_coordinate_text_report(planned: dict[str, Any]) -> dict[str, Any]:
    font = planned["font"]
    paragraph = planned.get("paragraph", {})
    return {
        "ownership_id": planned["ownership_id"],
        "split_unit_id": planned.get("split_unit_id"),
        "style_profile_id": planned.get("style_profile_id"),
        "visual_group_id": planned.get("visual_group_id"),
        "text": planned["text"],
        "display_text": planned.get("display_text", planned["text"]),
        "wrap_policy": planned.get("wrap_policy", {}),
        "relative_box": planned["relative_box"],
        "box_px": planned["box_px"],
        "font": {
            "resolved_family": font["resolved_family"],
            "target_font_size_pt": font["target_font_size_pt"],
            "resolved_color": font["resolved_color"],
            "bold": font.get("bold"),
        },
        "paragraph": {
            "horizontal_align": paragraph.get("horizontal_align"),
            "vertical_align": paragraph.get("vertical_align"),
            "line_spacing": paragraph.get("line_spacing"),
            "text_box_insets_pt": paragraph.get("text_box_insets_pt"),
        },
    }


def _coordinate_slide_summary(
    plan_slide: dict[str, Any],
    unit_reports: list[dict[str, Any]],
    merged_candidates: list[dict[str, Any]],
) -> dict[str, Any]:
    planned_units = plan_slide.get("text_units", [])
    style_profile_ids = {
        unit.get("style_profile_id")
        for unit in planned_units
        if isinstance(unit.get("style_profile_id"), str) and unit["style_profile_id"].strip()
    }
    return {
        "expected_text_units": len(planned_units),
        "expected_native_elements": len(plan_slide.get("native_elements", [])),
        "coordinate_text_units": len(unit_reports),
        "matched_text_shapes": sum(1 for item in unit_reports if item.get("actual")),
        "missing_text_units": sum(1 for item in unit_reports if not item.get("actual")),
        "possible_merged_text_shapes": len(merged_candidates),
        "deviations": sum(1 for item in unit_reports if not item["within_allowed_range"]),
        "style_profile_count": len(style_profile_ids),
        "units_with_split_unit_id": sum(1 for unit in planned_units if unit.get("split_unit_id")),
        "units_with_style_profile_id": sum(1 for unit in planned_units if unit.get("style_profile_id")),
    }


def _shape_relative_box(shape: Any, slide_width: int, slide_height: int) -> dict[str, float]:
    return {
        "x": round(int(shape.left) / slide_width, 6),
        "y": round(int(shape.top) / slide_height, 6),
        "w": round(int(shape.width) / slide_width, 6),
        "h": round(int(shape.height) / slide_height, 6),
    }


def _shape_font(shape: Any) -> dict[str, Any]:
    font = _first_run_font(shape)
    if font is None:
        return {}
    result: dict[str, Any] = {}
    if font.name:
        result["resolved_family"] = font.name
    if font.size is not None:
        result["target_font_size_pt"] = round(float(font.size.pt), 2)
    color = _font_color(font)
    if color:
        result["resolved_color"] = color
    return result


def _first_run_font(shape: Any) -> Any | None:
    if not getattr(shape, "has_text_frame", False):
        return None
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            return run.font
    return None


def _font_color(font: Any) -> str | None:
    try:
        rgb = font.color.rgb
    except AttributeError:
        return None
    if rgb is None:
        return None
    return f"#{str(rgb).upper()}"


def _shape_text(shape: Any) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return "\n".join(paragraph.text for paragraph in shape.text_frame.paragraphs)


def _possible_merged_text_shapes(slide: Any, text_units: list[dict[str, Any]]) -> list[dict[str, Any]]:
    planned = [(unit["text_unit_id"], unit["text"].strip()) for unit in text_units if unit.get("text", "").strip()]
    findings: list[dict[str, Any]] = []
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        text = _shape_text(shape).strip()
        if not text:
            continue
        matched = [text_unit_id for text_unit_id, planned_text in planned if planned_text and planned_text in text]
        if len(matched) > 1:
            findings.append(
                {
                    "shape_name": getattr(shape, "name", ""),
                    "matched_text_unit_ids": matched,
                    "reason": "one PPT text shape appears to contain multiple planned text units",
                }
            )
    return findings


def _text_within_plan(planned: dict[str, Any], actual: dict[str, Any]) -> tuple[bool, str]:
    actual_font = actual.get("font", {})
    planned_font = planned.get("font", {})
    required_font_fields = ("resolved_family", "target_font_size_pt", "resolved_color")
    missing = [field for field in required_font_fields if field not in actual_font]
    if missing:
        return False, "PPTX inspect could not read actual font fields: " + ", ".join(missing)
    allowed_range = planned_font["allowed_font_size_range_pt"]
    size = actual_font["target_font_size_pt"]
    if size < allowed_range[0] or size > allowed_range[1]:
        return False, "actual font size is outside allowed range"
    for field in ("resolved_family", "resolved_color"):
        if actual_font[field] != planned_font[field]:
            return False, f"actual {field} differs from plan"
    if not _relative_box_close(planned["relative_box"], actual["relative_box"]):
        return False, "actual relative_box differs from plan"
    return True, ""


def _relative_box_close(planned: dict[str, Any], actual: dict[str, Any], tolerance: float = 0.01) -> bool:
    for field in ("x", "y", "w", "h"):
        if abs(float(planned[field]) - float(actual[field])) > tolerance:
            return False
    return True


def _relative_box_distance(planned: dict[str, Any], actual: dict[str, Any]) -> float:
    return sum(abs(float(planned[field]) - float(actual[field])) for field in ("x", "y", "w", "h"))


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    return f"sha256:{digest}"
